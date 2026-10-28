from pathlib import Path
import json, re, sys
from docx import Document

root=Path(__file__).resolve().parents[1]
docx=root/'docs/report/BAOCAO_NOSQL_V2.docx'
required=['00_outline_old.md','01_reverify.md','02_stale_refs.md','STRUCTURE.md','FIGURES_TODO.md','FORMAT_NOTES.md','TRACE.md','SOURCE_PACKAGE.md','QA_PREP.md','SLIDE_OUTLINE.md','CHANGELOG.md','QUESTIONS_FOR_GROUP.md']
errors=[]
for name in required:
    if not (root/'docs/report'/name).is_file(): errors.append('missing docs/report/'+name)
if not docx.is_file(): errors.append('missing DOCX')
else:
    d=Document(docx)
    text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
    forbidden=['phase0','D:\\','shopquan_ao','_phase06_','phase07_','phase03_','WEBTHOITRANG','61 document','61 key','27 versus 21','27 vs 21','21 results','27 results','chưa kiểm thử']
    for term in forbidden:
        if term.casefold() in text.casefold(): errors.append('forbidden string: '+term)
    f=json.loads((root/'docs/phase1/phase1_facts.json').read_text(encoding='utf-8'))
    for n in [f['total_documents'],f['business_documents'],f['design_documents'],f['documents_by_type']['order'],f['guest_index_before_after']['after']['results']]:
        if not re.search(r'(?<!\d)'+str(n)+r'(?!\d)',text): errors.append('key count missing from report: '+str(n))
    normal=d.styles['Normal']
    if normal.font.name!='Times New Roman' or normal.font.size is None or round(normal.font.size.pt)!=13: errors.append('Normal style must be TNR 13pt')
    if normal.paragraph_format.line_spacing!=1.5: errors.append('Normal line spacing must be 1.5')
    other_fonts=[]
    for p in d.paragraphs:
        for r in p.runs:
            if r.font.name and r.font.name!='Times New Roman': other_fonts.append(r.font.name)
    if other_fonts: errors.append('non-TNR runs in body: '+str(len(other_fonts)))
    if len(d.inline_shapes)<7: errors.append('expected retained PHP UI images not embedded')
print('REPORT CHECK:', 'FAIL' if errors else 'PASS')
print('DOCX paragraphs=',len(d.paragraphs),'tables=',len(d.tables),'images=',len(d.inline_shapes))
print('errors=',json.dumps(errors,ensure_ascii=False))
sys.exit(1 if errors else 0)
