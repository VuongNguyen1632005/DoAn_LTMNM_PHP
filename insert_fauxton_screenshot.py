from pathlib import Path
import shutil

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement


report_path = Path(r"D:\PHP_NOSQL\BAOCAO_NOSQL_HOANCHINH.docx")
image_path = Path(r"D:\PHP_NOSQL\fauxton_order_statistics.png")
backup_path = report_path.with_name(report_path.stem + "_before_fauxton.docx")
caption = "Hình 3.14 – Kết quả chạy View order_count_by_status trong Fauxton"

if not report_path.exists():
    raise FileNotFoundError(report_path)
if not image_path.exists():
    raise FileNotFoundError(image_path)

doc = Document(str(report_path))

if any(p.text.strip() == caption for p in doc.paragraphs):
    print("Caption already exists; no change made.")
    raise SystemExit(0)

start_index = next(
    (i for i, p in enumerate(doc.paragraphs)
     if p.text.strip().startswith("3.4.13 Truy vấn thống kê nâng cao bằng MapReduce View")),
    None,
)
if start_index is None:
    raise RuntimeError("Could not locate section 3.4.13")

heading = next(
    (p for p in doc.paragraphs[start_index + 1:]
     if p.text.strip().startswith("3.5 ")),
    None,
)
if heading is None:
    raise RuntimeError("Could not locate the following section 3.5 heading")

shutil.copy2(report_path, backup_path)

image_xml = OxmlElement("w:p")
heading._p.addprevious(image_xml)
image_para = Paragraph(image_xml, heading._parent)
image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
image_run = image_para.add_run()
image_run.add_picture(str(image_path), width=Inches(6.1))

caption_xml = OxmlElement("w:p")
heading._p.addprevious(caption_xml)
caption_para = Paragraph(caption_xml, heading._parent)
caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
try:
    caption_para.style = "Caption"
except KeyError:
    caption_para.style = "Normal"
caption_run = caption_para.add_run(caption)
caption_run.font.size = Pt(10)

doc.save(str(report_path))
print(f"Updated: {report_path}")
print(f"Backup: {backup_path}")
