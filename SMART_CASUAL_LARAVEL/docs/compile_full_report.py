import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from pathlib import Path

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="B0C4DE"):
    tblPr = table._element.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

doc = docx.Document()

# Page setup: A4, Margins: Top 2cm, Bottom 2cm, Left 2.5cm, Right 2cm
for section in doc.sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.79)
    section.bottom_margin = Inches(0.79)
    section.left_margin = Inches(0.98)
    section.right_margin = Inches(0.79)

normal_style = doc.styles['Normal']
normal_style.font.name = 'Times New Roman'
normal_style.font.size = Pt(13)
normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
normal_style.paragraph_format.line_spacing = 1.3
normal_style.paragraph_format.space_after = Pt(4)

# ----------------- COVER PAGE -----------------
header_tbl = doc.add_table(rows=1, cols=2)
header_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
header_tbl.autofit = False

c0 = header_tbl.rows[0].cells[0]
c0.width = Inches(3.2)
p0 = c0.paragraphs[0]
p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p0.add_run("TRƯỜNG ĐH CÔNG THƯƠNG TP. HCM\n")
r.font.bold = True
r.font.size = Pt(11)
r2 = p0.add_run("KHOA CÔNG NGHỆ THÔNG TIN\nBỘ MÔN KỸ THUẬT PHẦN MỀM\n❖❖❖")
r2.font.size = Pt(10.5)

c1 = header_tbl.rows[0].cells[1]
c1.width = Inches(3.5)
p1 = c1.paragraphs[0]
p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p1.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
r.font.bold = True
r.font.size = Pt(11)
r2 = p1.add_run("Độc lập – Tự do – Hạnh phúc\n")
r2.font.bold = True
r2.font.size = Pt(10.5)
r3 = p1.add_run("— — — — — — — — —")
r3.font.size = Pt(10)

doc.add_paragraph().paragraph_format.space_after = Pt(28)

title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_t1 = title_p.add_run("BÁO CÁO ĐỒ ÁN CUỐI MÔN HỌC\n")
r_t1.font.bold = True
r_t1.font.size = Pt(15)
r_t1.font.color.rgb = RGBColor(0x14, 0x22, 0x1D)

r_t2 = title_p.add_run("HỌC PHẦN: LẬP TRÌNH MÃ NGUỒN MỞ\n\n")
r_t2.font.bold = True
r_t2.font.size = Pt(13)

r_t3 = title_p.add_run("ĐỀ TÀI:\nXÂY DỰNG WEBSITE BÁN HÀNG THỜI TRANG \"SMART CASUAL\"\nBẰNG LARAVEL FRAMEWORK VÀ CƠ SỞ DỮ LIỆU MYSQL\n")
r_t3.font.bold = True
r_t3.font.size = Pt(16)
r_t3.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

doc.add_paragraph().paragraph_format.space_after = Pt(36)

# Information table
info_tbl = doc.add_table(rows=6, cols=2)
info_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(info_tbl, "14221D")

info_data = [
    ("Sinh viên 1 (Trưởng nhóm):", "Nguyễn Huy Hoàng — MSSV: 2001230266 — Lớp: 14DHTH02"),
    ("Sinh viên 2 (Thành viên):", "Nguyễn Việt Vương — MSSV: 2001231062 — Lớp: 14DHTH02"),
    ("Sinh viên 3 (Thành viên):", "Nguyễn Ngọc Tường Vi — MSSV: 2001231032 — Lớp: 14DHTH12"),
    ("Giảng viên hướng dẫn:", "Thầy Nguyễn Minh Hiến"),
    ("Ngành đào tạo:", "Công nghệ Thông tin — Chuyên ngành Kỹ thuật Phần mềm"),
    ("Thời gian thực hiện:", "Học kỳ 1 — Năm học 2026 – 2027")
]

for idx, (label, val) in enumerate(info_data):
    row = info_tbl.rows[idx]
    cell0 = row.cells[0]
    cell0.width = Inches(2.2)
    set_cell_background(cell0, "F4F6F8")
    set_cell_margins(cell0, 70, 70, 100, 100)
    p = cell0.paragraphs[0]
    r = p.add_run(label)
    r.font.bold = True
    r.font.size = Pt(11)
    
    cell1 = row.cells[1]
    cell1.width = Inches(4.5)
    set_cell_margins(cell1, 70, 70, 100, 100)
    p = cell1.paragraphs[0]
    r = p.add_run(val)
    r.font.size = Pt(11)

doc.add_page_break()

# ----------------- TABLE OF CONTENTS -----------------
h_toc = doc.add_paragraph()
h_toc.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h_toc.add_run("MỤC LỤC BÁO CÁO\n")
r.font.bold = True
r.font.size = Pt(15)

toc_items = [
    ("1. GIỚI THIỆU ĐỀ TÀI", True),
    ("   1.1. Lý do chọn đề tài", False),
    ("   1.2. Mục tiêu đề tài", False),
    ("   1.3. Phạm vi thực hiện", False),
    ("2. PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG", True),
    ("   2.1. Phân tích đối tượng và vai trò người dùng", False),
    ("   2.2. Sơ đồ UML Use Case tổng quan", False),
    ("   2.3. Sơ đồ và thiết kế Cơ sở dữ liệu (ERD)", False),
    ("   2.4. Danh mục 13 bảng nghiệp vụ chi tiết", False),
    ("   2.5. Thiết kế Triggers và tính toán tự động", False),
    ("3. CÔNG NGHỆ VÀ KIẾN TRÚC TRIỂN KHAI", True),
    ("   3.1. Kiến trúc tổng thể hệ thống", False),
    ("   3.2. Danh mục công nghệ và thư viện sử dụng", False),
    ("   3.3. Cơ chế xác thực và phân quyền (Multi-Auth & Middleware)", False),
    ("   3.4. Mô hình ORM Eloquent và bảo vệ toàn vẹn dữ liệu", False),
    ("4. HƯỚNG DẪN CÀI ĐẶT VÀ SỬ DỤNG", True),
    ("   4.1. Yêu cầu môi trường hệ thống", False),
    ("   4.2. Các bước cài đặt chi tiết từ mã nguồn", False),
    ("   4.3. Lỗi thường gặp khi cài đặt và cách xử lý", False),
    ("   4.4. Danh sách tài khoản thử nghiệm hệ thống", False),
    ("5. KẾT QUẢ ĐẠT ĐƯỢC VÀ CHỨC NĂNG HỆ THỐNG", True),
    ("   5.1. Phân hệ Khách hàng (Customer Storefront)", False),
    ("   5.2. Phân hệ Quản trị viên (Admin Portal)", False),
    ("   5.3. Bảng đối chiếu yêu cầu kỹ thuật (Traceability Matrix)", False),
    ("6. KIỂM THỬ VÀ ĐÁNH GIÁ CHẤT LƯỢNG SẢN PHẨM", True),
    ("   6.1. Phương pháp kiểm thử tự động (Feature Tests)", False),
    ("   6.2. Kết quả kiểm thử 23 kịch bản kiểm thử tự động", False),
    ("7. PHÂN CÔNG CÔNG VIỆC VÀ ĐÁNH GIÁ ĐÓNG GÓP NHÓM", True),
    ("   7.1. Bảng phân công nhiệm vụ chi tiết", False),
    ("   7.2. Tỷ lệ hoàn thành và phương thức phối hợp nhóm", False),
    ("8. HẠN CHẾ VÀ HƯỚNG PHÁT TRIỂN", True),
    ("   8.1. Các hạn chế hiện tại", False),
    ("   8.2. Kế hoạch nâng cấp và phát triển mở rộng", False),
    ("9. KẾT LUẬN", True),
    ("10. TÀI LIỆU THAM KHẢO", True),
]

for title, is_bold in toc_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(title)
    r.font.bold = is_bold
    r.font.size = Pt(11.5 if is_bold else 11)

doc.add_page_break()

# ----------------- BODY HELPERS -----------------
def add_chapter(num_title):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(num_title)
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x14, 0x22, 0x1D)
    return h

def add_sec(title):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(title)
    r.font.bold = True
    r.font.size = Pt(12.5)
    r.font.color.rgb = RGBColor(0x24, 0x3B, 0x32)
    return h

def add_para(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run(text)
    r.font.size = Pt(12.5)
    return p

def add_bullet(bold_prefix, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    r1 = p.add_run("• " + bold_prefix)
    r1.font.bold = True
    r1.font.size = Pt(12)
    r2 = p.add_run(text)
    r2.font.size = Pt(12)
    return p

def add_caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

def add_placeholder_box(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"[{text}]")
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

# ----------------- CHƯƠNG 1 -----------------
add_chapter("1. GIỚI THIỆU ĐỀ TÀI")
add_sec("1.1. Lý do chọn đề tài")
add_para("Các cửa hàng kinh doanh thời trang bán lẻ quy mô vừa và nhỏ thường gặp khó khăn trong việc kiểm soát số lượng tồn kho theo từng kích cỡ sản phẩm. Một mẫu áo hay một mẫu giày khi nhập về luôn có nhiều size khác nhau, chẳng hạn áo có size S, M, L và giày có size 39, 40, 41. Mỗi kích cỡ có lượng hàng trong kho riêng và đôi khi có mức giá bán khác nhau. Khi vận hành bằng sổ sách hoặc nhận đơn thủ công qua tin nhắn mạng xã hội, nhân viên bán hàng rất dễ nhận đơn cho những kích cỡ đã hết hàng nhưng chưa kịp kiểm đếm thực tế. Việc thiếu đồng bộ giữa dữ liệu trưng bày và dữ liệu tồn kho thực dẫn đến tình trạng hủy đơn, làm giảm uy tín của cửa hàng.")
add_para("Bên cạnh vấn đề tồn kho, quy trình xử lý đơn hàng từ lúc khách đặt mua, nhân viên xác nhận, đóng gói cho đến lúc giao vận và thu tiền đòi hỏi sự phân định trạng thái rõ ràng. Nếu các trạng thái này không được quản lý bằng một hệ thống tập trung, nhân sự phụ trách kho và nhân viên bán hàng sẽ mất nhiều thời gian trao đổi nội bộ, dẫn đến việc chậm trễ tiến độ giao hàng và khó khăn trong đối soát tài chính vào cuối tháng.")
add_para("Từ những vấn đề thực tế trên, nhóm chọn đề tài \"Xây dựng Website Bán hàng Thời trang Smart Casual bằng Laravel Framework và MySQL\". Đề tài tập trung giải quyết bài toán quản lý danh mục sản phẩm có phân tách biến thể kích thước, giỏ hàng trực tuyến và quy trình đặt hàng kèm trừ tồn kho tự động. Thông qua đồ án này, nhóm có cơ hội đưa các kiến thức đã học trong học phần Lập trình mã nguồn mở, bao gồm lập trình PHP hướng đối tượng, kiến trúc MVC của Laravel Framework, thao tác CSDL quan hệ với Eloquent ORM và phân quyền người dùng bằng Middleware, vào giải quyết một bài toán cụ thể.")

add_sec("1.2. Mục tiêu đề tài")
add_para("Về chuẩn đầu ra học phần (CLOs):")
add_bullet("CLO1: ", "Vận dụng PHP 8.3, kiến trúc hướng đối tượng, cơ sở dữ liệu quan hệ MySQL và framework Laravel 11.x để xây dựng hoàn chỉnh ứng dụng web thương mại điện tử.")
add_bullet("CLO2: ", "Thiết kế giao diện bằng Blade Template kết hợp Bootstrap 5.2.3, đảm bảo khả năng hiển thị tương thích trên các kích thước màn hình máy tính và thiết bị di động, cung cấp phản hồi thao tác rõ ràng cho người dùng.")
add_bullet("CLO3: ", "Phân chia công việc cụ thể cho từng thành viên, phối hợp quản lý mã nguồn và theo dõi tiến độ thực hiện thông qua hệ thống Git và GitHub.")
add_bullet("CLO4: ", "Xây dựng phần mềm hoàn chỉnh có khả năng đưa vào sử dụng thực tế cho cửa hàng thời trang với đầy đủ luồng xem hàng, giỏ hàng, thanh toán, xử lý đơn và thống kê doanh thu.")
add_bullet("CLO5: ", "Thể hiện tinh thần tự học, tìm hiểu sâu các cơ chế xử lý tranh chấp dữ liệu, bảo mật phân quyền và tuân thủ các quy định về liêm chính học thuật.")

add_para("Về mục tiêu kỹ thuật phần mềm:")
add_bullet("Cơ sở dữ liệu: ", "Thiết kế cơ sở dữ liệu quan hệ gồm 13 bảng ở dạng chuẩn 3NF, thiết lập đầy đủ ràng buộc khóa ngoại và xây dựng 4 Database Triggers nhằm tự động hóa việc tính toán tổng số lượng và tiền hàng.")
add_bullet("Kiến trúc mã nguồn: ", "Phân tách mã nguồn rõ ràng theo mô hình MVC kết hợp Service Layer: Controller làm nhiệm vụ điều phối HTTP request, Form Request xử lý xác thực dữ liệu đầu vào, Service Layer đảm nhiệm logic thanh toán và giỏ hàng, Model Eloquent chịu trách nhiệm ánh xạ và truy vấn dữ liệu.")
add_bullet("Bảo mật và an toàn dữ liệu: ", "Xây dựng hệ thống bảo mật gồm: cơ chế Multi-Auth phân tách bảo vệ giữa khách hàng và quản trị viên, kiểm soát phân quyền qua Middleware, ngăn chặn truy cập trái phép vào đơn hàng của người khác (IDOR), băm mật khẩu bằng thuật toán Bcrypt và chống tấn công CSRF trên các form gửi dữ liệu.")
add_bullet("Kiểm thử chất lượng: ", "Kiểm chứng độ tin cậy của mã nguồn bằng bộ kiểm thử tự động (Feature Tests) gồm 23 kịch bản kiểm thử, bao phủ các trường hợp xử lý tranh chấp tồn kho, phân quyền và tính toán hóa đơn.")

add_sec("1.3. Phạm vi thực hiện")
add_para("Đối tượng phục vụ của hệ thống gồm: Khách hàng vãng lai (Guest) xem catalog và đặt hàng; Khách hàng thành viên (Customer) quản lý tài khoản, lịch sử đơn và viết đánh giá; Nhân viên bán hàng (Staff) tiếp nhận đơn, cập nhật giao vận và xác nhận thanh toán; Quản trị viên (Manager) toàn quyền quản lý danh mục, sản phẩm và xem báo cáo tài chính.")
add_para("Phạm vi chức năng bao gồm: Quản lý danh mục và kích cỡ; Quản lý sản phẩm, biến thể tồn kho/giá và upload hình ảnh qua Laravel Storage; Tìm kiếm và lọc đa tiêu chí kèm phân trang giữ tham số truy vấn; Giỏ hàng phiên làm việc và giỏ hàng CSDL tự động đồng bộ khi đăng nhập; Thanh toán nguyên tử trong Database Transaction có khóa dòng dữ liệu chống bán vượt tồn kho; Đánh giá nhận xét có kiểm duyệt; Báo cáo doanh thu tài chính theo tháng.")

# ----------------- CHƯƠNG 2 -----------------
add_chapter("2. PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG")
add_sec("2.1. Phân tích đối tượng và vai trò người dùng")
add_para("Trong hệ thống Smart Casual, nhóm phân tách người dùng thành các đối tượng độc lập thay vì gom chung vào một bảng người dùng duy nhất:")
add_bullet("Khách hàng vãng lai (Guest): ", "Người dùng chưa đăng nhập, có thể duyệt catalog, tìm kiếm, xem chi tiết và thêm hàng vào giỏ hàng phiên làm việc (Session). Cho phép khách vãng lai đặt hàng trực tiếp mà không bắt buộc tạo tài khoản.")
add_bullet("Khách hàng thành viên (Customer): ", "Khách hàng đăng ký tài khoản bằng email. Khi đăng nhập, toàn bộ giỏ hàng Session được chuyển vào bảng GIO_HANG. Khách hàng theo dõi lịch sử đơn cá nhân tại /account/orders, hủy đơn đang chờ duyệt, xác nhận đã nhận hàng và gửi đánh giá sản phẩm. Khách hàng bị chặn hoàn toàn khi cố truy cập /admin.")
add_bullet("Nhân viên bán hàng (Staff): ", "Nhân sự xử lý đơn hàng hằng ngày, đăng nhập qua /admin/login. Nhân viên có quyền xem danh sách đơn hàng, cập nhật các giai đoạn giao nhận, xác nhận tình trạng đã thu tiền và duyệt đánh giá. Nhân viên không được phép sửa giá sản phẩm, không được xóa danh mục và không được xem báo cáo doanh thu.")
add_bullet("Quản trị viên (Manager): ", "Người phụ trách kinh doanh cao nhất. Manager có toàn bộ quyền của Staff, đồng thời có độc quyền truy cập vào các route CRUD danh mục, CRUD sản phẩm, ẩn sản phẩm và xem báo cáo doanh thu tài chính.")

add_sec("2.2. Sơ đồ UML Use Case tổng quan")
add_para("Sơ đồ Use Case tổng quan mô tả mối quan hệ giữa 4 nhóm tác nhân và các chức năng chính của hệ thống:")
add_placeholder_box("SƠ ĐỒ USE CASE HỆ THỐNG SMART CASUAL")
add_caption("Hình 2.1: Sơ đồ Use Case tổng quan thể hiện tương tác của từng vai trò trong hệ thống Smart Casual.")

add_sec("2.3. Sơ đồ và thiết kế Cơ sở dữ liệu (ERD)")
add_para("Cơ sở dữ liệu được xây dựng trên MySQL 8.4 gồm 13 bảng nghiệp vụ. Nhóm đã đưa ra các quyết định thiết kế then chốt:")
add_bullet("Tách riêng SAN_PHAM và CHI_TIET_SP: ", "Một sản phẩm có nhiều kích cỡ (S, M, L hoặc 39, 40, 41). Mỗi kích cỡ có lượng tồn kho riêng và đôi khi có đơn giá bán khác nhau. Việc tách bảng biến thể CHI_TIET_SP cho phép hệ thống kiểm soát chính xác lượng tồn kho của từng size, ngăn chặn việc đặt mua khi size đó đã hết hàng.")
add_bullet("Lưu đơn giá snapshot trong CHI_TIET_DON_HANG: ", "Khi khách đặt hàng, đơn giá được sao chép cố định vào cột DonGia của chi tiết đơn. Nếu giá bán của sản phẩm trong CHI_TIET_SP thay đổi trong tương lai, tổng tiền của các hóa đơn cũ trong quá khứ vẫn được bảo toàn nguyên vẹn.")
add_bullet("Sử dụng khóa chính phức hợp: ", "Bảng GIO_HANG dùng cặp (MaKH, MaCTSP) và bảng CHI_TIET_DON_HANG dùng cặp (MaDonHang, MaCTSP) làm khóa chính, đảm bảo mỗi biến thể chỉ xuất hiện một lần trong giỏ hàng hoặc đơn hàng.")

add_placeholder_box("SƠ ĐỒ QUAN HỆ THỰC THỂ (ERD) 13 BẢNG NGHIỆP VỤ")
add_caption("Hình 2.2: Sơ đồ quan hệ thực thể (ERD) của hệ thống thương mại điện tử Smart Casual.")

add_sec("2.4. Danh mục 13 bảng nghiệp vụ chi tiết")
add_para("Dưới đây là cấu trúc 13 bảng nghiệp vụ được khởi tạo qua Laravel Migrations:")

table_schemas = [
    ("1. DANH_MUC (Danh mục sản phẩm)", "MaDanhMuc (VARCHAR(10), PK), TenDanhMuc (VARCHAR(100)), TrangThai (BOOLEAN)", "Bảng phân loại hàng hóa chính (Áo, Quần, Giày), dùng làm tiêu chí lọc và phân cấp danh mục."),
    ("2. KICH_THUOC (Bảng kích cỡ chuẩn)", "MaKichThuoc (VARCHAR(10), PK), TenKichThuoc (VARCHAR(50)), TrangThai (BOOLEAN)", "Quy chuẩn hóa danh sách kích cỡ dùng chung (S, M, L, XL, 38, 39, 40...), tránh nhập tự do sai chính tả."),
    ("3. SAN_PHAM (Sản phẩm gốc)", "MaSP (VARCHAR(20), PK), MaDanhMuc (VARCHAR(10), FK), TenSanPham (VARCHAR(200)), MoTa (TEXT), TrangThai (BOOLEAN)", "Lưu thông tin chung của sản phẩm. Trường TrangThai phục vụ xóa mềm logic để bảo vệ khóa ngoại với đơn hàng cũ."),
    ("4. CHI_TIET_SP (Biến thể chi tiết theo Size)", "MaCTSP (VARCHAR(30), PK), MaSP (VARCHAR(20), FK), MaKichThuoc (VARCHAR(10), FK), GiaBan (DECIMAL(18,2)), SoLuong (INT UNSIGNED), TrangThai (BOOLEAN). Ràng buộc UNIQUE(MaSP, MaKichThuoc)", "Đơn vị lưu kho SKU thực tế, là trọng tâm trong các truy vấn khóa dòng và trừ tồn kho khi checkout."),
    ("5. HINH_ANH_SP (Bộ sưu tập ảnh sản phẩm)", "MaHinhAnh (VARCHAR(30), PK), MaSP (VARCHAR(20), FK), DuongDan (VARCHAR(255)), LaAnhChinh (BOOLEAN), TrangThai (BOOLEAN)", "Quản lý nhiều ảnh chụp cho mỗi sản phẩm mà không làm cồng kềnh bảng SAN_PHAM."),
    ("6. KHACH_HANG (Tài khoản khách hàng)", "MaKH (VARCHAR(20), PK), TenDangNhap (VARCHAR(100), UK), MatKhau (VARCHAR(255)), HoTen (VARCHAR(100)), SoDienThoai (VARCHAR(20)), Email (VARCHAR(100), UK), DiaChi (VARCHAR(255)), TrangThai (BOOLEAN)", "Phục vụ riêng cho Guard web. Khóa chính là chuỗi ký tự MaKH thay vì số tự tăng id."),
    ("7. NHAN_VIEN (Tài khoản nhân sự và quản trị)", "MaNV (VARCHAR(20), PK), TenDangNhap (VARCHAR(100), UK), MatKhau (VARCHAR(255)), HoTen (VARCHAR(100)), ChucVu (VARCHAR(50): manager / staff), TrangThai (BOOLEAN)", "Phục vụ Guard admin. Phân tách rạch ròi với bảng KHACH_HANG để ngăn ngừa phân quyền nhầm lẫn."),
    ("8. MA_GIAM_GIA (Mã khuyến mãi & Voucher)", "MaCode (VARCHAR(50), PK), GiaTri (DECIMAL(18,2)), LoaiGiamGia (VARCHAR(20): percent / cash / shipping), SoLuong (INT), TrangThai (BOOLEAN)", "Cột SoLuong bị trừ 1 khi đặt hàng thành công và được hoàn trả khi đơn hàng bị hủy trước khâu đóng gói."),
    ("9. PHUONG_THUC_VAN_CHUYEN (Đơn vị giao hàng)", "MaPTVC (VARCHAR(20), PK), TenPTVC (VARCHAR(100)), GiaVanChuyen (DECIMAL(18,2)), TrangThai (BOOLEAN)", "Quy định cước phí cố định ('HT': 30.000đ, 'TC': 15.000đ) để áp tự động vào hóa đơn."),
    ("10. GIO_HANG (Giỏ hàng thành viên lưu CSDL)", "MaKH (VARCHAR(20), FK), MaCTSP (VARCHAR(30), FK), SoLuong (INT). Khóa chính: PRIMARY KEY(MaKH, MaCTSP)", "Sử dụng khóa chính phức hợp hai cột, đòi hỏi ghi đè setKeysForSaveQuery trong Eloquent Model."),
    ("11. DON_HANG (Hóa đơn và thông tin đặt hàng)", "MaDonHang (VARCHAR(30), PK), MaKH (FK, NULL), MaNV (FK, NULL), MaCode (FK, NULL), MaPTVC (FK), TenNguoiNhan (VARCHAR(100)), SoDienThoai (VARCHAR(20)), Email (VARCHAR(100)), NgayDat (DATETIME), DiaChiGiao (VARCHAR(255)), TongSoLuong (INT), TongTien (DECIMAL(18,2)), GiamGia (DECIMAL(18,2)), PhiVanChuyen (DECIMAL(18,2)), ThanhTien (DECIMAL(18,2)), PhuongThucThanhToan (VARCHAR(50)), DaThanhToan (BOOLEAN), TrangThaiDonHang (VARCHAR(50))", "Lưu thông tin người nhận tại thời điểm đặt, tránh việc đổi địa chỉ tài khoản làm sai lệch đơn hàng cũ."),
    ("12. CHI_TIET_DON_HANG (Chi tiết món trong đơn)", "MaDonHang (VARCHAR(30), FK), MaCTSP (VARCHAR(30), FK), SoLuong (INT), DonGia (DECIMAL(18,2)), ThanhTien (DECIMAL(18,2)). Khóa chính: PRIMARY KEY(MaDonHang, MaCTSP)", "Quy tắc ON DELETE RESTRICT trên MaCTSP ngăn chặn xóa biến thể khi đã phát sinh đơn hàng lịch sử."),
    ("13. DANH_GIA (Đánh giá nhận xét của khách)", "MaDanhGia (VARCHAR(30), PK), MaSP (VARCHAR(20), FK), MaKH (VARCHAR(20), FK, NULL), TenKhachHang (VARCHAR(100)), SoSao (TINYINT), NoiDung (TEXT), TrangThai (BOOLEAN)", "Trường TrangThai giúp nhân viên kiểm duyệt ẩn các nhận xét vi phạm hoặc spam trước khi công khai.")
]

for t_name, t_cols, t_note in table_schemas:
    add_bullet(t_name + ": ", f"Cột: {t_cols}. {t_note}")

add_sec("2.5. Thiết kế Triggers và tính toán tự động")
add_para("Nhóm phân chia trách nhiệm tính toán tài chính và quản lý đơn hàng thành hai tầng:")
add_bullet("Tầng Database Triggers (MySQL Triggers): ", "Gồm 4 triggers: trg_tudong_lay_phiship_insert tự động lấy cước phí vận chuyển và tính ThanhTien = GREATEST(0, TongTien + PhiVanChuyen - GiamGia); các triggers AFTER INSERT, AFTER UPDATE, AFTER DELETE trên CHI_TIET_DON_HANG tự động tính lại TongSoLuong và TongTien bằng hàm SUM rồi cập nhật vào DON_HANG. Khi chi tiết đơn thay đổi, tổng tiền luôn tự động cập nhật mà không sợ lập trình viên quên viết code.")
add_bullet("Tầng Application Service (CheckoutService): ", "Thực thi các nghiệp vụ ứng dụng: kiểm tra tính hợp lệ của voucher qua isValid(); bọc chu trình đặt hàng trong DB::transaction(); khóa dòng dữ liệu biến thể bằng lockForUpdate() để ngăn chặn Race Condition; và kích hoạt quy trình hoàn kho, hoàn voucher khi đơn bị hủy theo ma trận ALLOWED_TRANSITIONS.")
add_para("Đánh đổi kỹ thuật: Cách làm này bảo toàn tính đúng đắn của dữ liệu ở mức lưu trữ nhưng khiến logic tính tiền nằm ở cả hai nơi (file migration tạo trigger và class CheckoutService), đòi hỏi phải đồng bộ cả hai khi có thay đổi về quy tắc tính phí.")

# ----------------- CHƯƠNG 3 -----------------
add_chapter("3. CÔNG NGHỆ VÀ KIẾN TRÚC TRIỂN KHAI")
add_sec("3.1. Kiến trúc tổng thể hệ thống")
add_para("Hệ thống áp dụng mô hình MVC mở rộng với lớp dịch vụ (Service Layer) của Laravel 11.x:")
add_para("Luồng xử lý của một request: public/index.php tiếp nhận request -> Router (routes/web.php) định tuyến -> Middleware Pipeline kiểm tra phiên và phân quyền (AdminMiddleware, ManagerOnlyMiddleware, throttle:login) -> Form Request xác thực dữ liệu đầu vào -> Controller ủy quyền cho Service Layer (CartService, CheckoutService) xử lý nghiệp vụ phức tạp -> Eloquent Models truy vấn xuống MySQL 8.4 qua PDO Parameter Binding an toàn -> Controller truyền kết quả sang View Blade để render giao diện HTML gửi về cho người dùng.")
add_placeholder_box("SƠ ĐỒ KIẾN TRÚC VÀ LUỒNG XỬ LÝ REQUEST CỦA HỆ THỐNG")
add_caption("Hình 3.1: Sơ đồ luồng xử lý HTTP request qua các tầng kiến trúc của hệ thống Smart Casual.")

add_sec("3.2. Danh mục công nghệ và thư viện sử dụng")
add_bullet("PHP 8.3: ", "Tận dụng typed properties, match expressions, constructor property promotion.")
add_bullet("Laravel 11.x: ", "Phiên bản ổn định tại thời điểm nhóm khởi tạo dự án, cung cấp cấu trúc MVC, Eloquent ORM, Routing và Migrations.")
add_bullet("MySQL 8.4 Community Server: ", "Hỗ trợ InnoDB, khóa ngoại, ACID transactions, lockForUpdate và Database Triggers.")
add_bullet("Blade & Bootstrap 5.2.3: ", "Template Engine tích hợp sẵn kết hợp lưới Bootstrap và FontAwesome 6.")
add_bullet("PHPUnit 10.5.66: ", "Thực thi bộ kiểm thử tự động trên cơ sở dữ liệu thử nghiệm độc lập.")
add_bullet("Composer 2.9.4: ", "Quản lý các gói phụ thuộc theo chuẩn PSR-4.")

add_sec("3.3. Cơ chế xác thực và phân quyền (Multi-Auth & Middleware)")
add_para("Hệ thống cấu hình hai Guard độc lập trong config/auth.php:")
add_bullet("Guard web: ", "Sử dụng provider khach_hangs trỏ vào Model KhachHang (bảng KHACH_HANG, khóa chính MaKH, mật khẩu MatKhau).")
add_bullet("Guard admin: ", "Sử dụng provider nhan_viens trỏ vào Model NhanVien (bảng NHAN_VIEN, khóa chính MaNV, mật khẩu MatKhau).")
add_para("Lý do tách hai Guard: Người mua hàng và nhân sự quản trị có các trường thông tin và vòng đời tài khoản hoàn toàn khác biệt. Việc tách riêng giúp một người dùng có thể đồng thời đăng nhập tài khoản khách hàng ở giao diện ngoài và đăng nhập tài khoản nhân viên ở giao diện quản trị trên cùng một trình duyệt mà không bị ghi đè Session.")
add_para("Phân quyền qua Middleware: AdminMiddleware chặn khách thường và khách vãng lai vào /admin; ManagerOnlyMiddleware chặn nhân viên staff vào báo cáo doanh thu và các thao tác thêm/sửa/xóa sản phẩm, danh mục. Route xem đơn hàng /account/orders/{id} được ràng buộc điều kiện MaKH = Auth::user()->MaKH, trả về HTTP 403 nếu cố tình xem đơn của người khác, ngăn chặn hoàn toàn lỗ hổng IDOR.")

add_sec("3.4. Mô hình ORM Eloquent và bảo vệ toàn vẹn dữ liệu")
add_bullet("Giải quyết Composite Primary Key: ", "Eloquent mặc định không hỗ trợ khóa chính phức hợp hai cột trên bảng GIO_HANG (MaKH, MaCTSP). Nhóm đã ghi đè phương thức setKeysForSaveQuery trong Model GioHang để ép buộc Eloquent đưa đủ cả hai điều kiện WHERE MaKH = ? AND MaCTSP = ? vào câu lệnh cập nhật, ngăn ngừa việc ghi đè lên các sản phẩm khác.")
add_bullet("Pessimistic Locking chống Race Condition: ", "Trong CheckoutService, nhóm sử dụng ChiTietSanPham::where('MaCTSP', $id)->lockForUpdate()->first() bên trong DB::transaction(). Khi có hai khách hàng cùng bấm đặt hàng cho sản phẩm chỉ còn tồn kho bằng 1, lệnh này khóa dòng dữ liệu tại mức CSDL InnoDB, buộc request thứ hai phải xếp hàng đợi cho đến khi transaction thứ nhất hoàn tất. Lúc request thứ hai đọc lại thì tồn kho đã bằng 0 và hệ thống sẽ báo lỗi hết hàng, giải quyết triệt để vấn đề Overselling.")
add_bullet("Bảo mật an toàn: ", "Chống SQL Injection qua PDO Parameter Binding của Eloquent; Chống CSRF qua token @csrf trên tất cả các form; Mã hóa đầu ra chống XSS qua cú pháp {{ }} của Blade.")

# ----------------- CHƯƠNG 4 -----------------
add_chapter("4. HƯỚNG DẪN CÀI ĐẶT VÀ SỬ DỤNG")
add_sec("4.1. Yêu cầu môi trường hệ thống")
add_para("Hệ thống yêu cầu: HĐH Windows 10/11, Linux hoặc macOS; PHP >= 8.2 (khuyên dùng PHP 8.3 với đầy đủ extension pdo_mysql, mbstring, openssl, tokenizer, xml, ctype, json, curl, fileinfo); Composer >= 2.2; MySQL 8.x chạy trên cổng 3306 mặc định.")

add_sec("4.2. Các bước cài đặt chi tiết từ mã nguồn")
add_para("Bước 1: Mở Terminal và di chuyển vào thư mục dự án: cd SMART_CASUAL_LARAVEL")
add_para("Bước 2: Cài đặt thư viện phụ thuộc: composer install --prefer-dist --no-interaction")
add_para("Bước 3: Thiết lập file môi trường: copy .env.example .env và cấu hình CSDL:")
add_para("DB_CONNECTION=mysql\nDB_HOST=127.0.0.1\nDB_PORT=3306\nDB_DATABASE=smart_casual_shop\nDB_USERNAME=root\nDB_PASSWORD=")
add_para("Bước 4: Sinh mã khóa ứng dụng và tạo symlink lưu trữ ảnh:\nphp artisan key:generate\nphp artisan storage:link")
add_para("Bước 5: Tạo database trong MySQL (CREATE DATABASE smart_casual_shop CHARACTER SET utf8mb4;) rồi chạy migration kết hợp seeder:\nphp artisan migrate:fresh --seed\n(Hoặc import trực tiếp file kết xuất SQL có sẵn tại database/smart_casual_shop_dump.sql).")
add_para("Bước 6: Khởi chạy máy chủ phát triển nội bộ: php artisan serve và truy cập tại http://localhost:8000")

add_sec("4.3. Lỗi thường gặp khi cài đặt và cách xử lý")
add_bullet("Lỗi No application encryption key has been specified: ", "Do file .env chưa có APP_KEY. Xử lý bằng cách chạy php artisan key:generate.")
add_bullet("Lỗi Unknown database 'smart_casual_shop': ", "Do chưa tạo database trước khi chạy migrate. Xử lý bằng cách vào phpMyAdmin hoặc MySQL CLI chạy lệnh CREATE DATABASE.")
add_bullet("Lỗi ảnh sản phẩm bị vỡ (HTTP 404): ", "Do chưa tạo liên kết lưu trữ. Xử lý bằng cách chạy php artisan storage:link với quyền quản trị.")
add_bullet("Lỗi Fixture file not found khi chạy seeder: ", "Do file fixture nạp ngoài không có sẵn. Xử lý bằng cách import trực tiếp file database/smart_casual_shop_dump.sql vào MySQL.")

add_sec("4.4. Danh sách tài khoản thử nghiệm hệ thống")
demo_tbl = doc.add_table(rows=5, cols=4)
demo_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(demo_tbl)

d_headers = ["Vai trò (Role)", "Cổng đăng nhập", "Tài khoản / Email", "Mật khẩu"]
for i, h in enumerate(d_headers):
    c = demo_tbl.rows[0].cells[i]
    set_cell_background(c, "14221D")
    set_cell_margins(c, 60, 60, 80, 80)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

demo_rows = [
    ("Khách hàng 1 (Customer)", "/login", "demo1@gmail.com", "password123"),
    ("Khách hàng 2 (IDOR Test)", "/login", "demo2@gmail.com", "password123"),
    ("Nhân viên (Staff)", "/admin/login", "staff", "staff123"),
    ("Quản trị (Manager)", "/admin/login", "admin", "admin123")
]

for r_idx, (r0, r1, r2, r3) in enumerate(demo_rows, start=1):
    row = demo_tbl.rows[r_idx]
    bg = "F9FBFA" if r_idx % 2 == 1 else "FFFFFF"
    for c_idx, val in enumerate([r0, r1, r2, r3]):
        c = row.cells[c_idx]
        set_cell_background(c, bg)
        set_cell_margins(c, 50, 50, 60, 60)
        p = c.paragraphs[0]
        if c_idx in [1, 3]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(val)
        r.font.size = Pt(10)

# ----------------- CHƯƠNG 5 -----------------
add_chapter("5. KẾT QUẢ ĐẠT ĐƯỢC VÀ CHỨC NĂNG HỆ THỐNG")
add_sec("5.1. Phân hệ Khách hàng (Customer Storefront)")
add_para("Quy trình mua sắm của khách hàng diễn ra theo luồng trải nghiệm thực tế:")
add_para("Khách hàng vào trang chủ (Route /) xem các sản phẩm thịnh hành và duyệt danh mục 48 sản phẩm thời trang. Khách sử dụng bộ lọc tìm kiếm theo tên, danh mục (Áo, Quần, Giày), khoảng giá và sắp xếp tăng/giảm. Hệ thống phân trang 16 sản phẩm/trang và giữ lại tham số lọc trên URL qua withQueryString().")
add_placeholder_box("HÌNH 5.1: Giao diện trang chủ và bộ lọc catalog sản phẩm tại route /")

add_para("Khách bấm vào sản phẩm để sang trang Chi tiết (Route /products/{id}), xem ảnh thumbnail chuyển đổi bằng JavaScript, bấm chọn kích cỡ (Size) để cập nhật giá và tồn kho khả dụng tương ứng. Khách xem điểm đánh giá sao, đọc nhận xét cũ và có thể gửi đánh giá mới.")
add_placeholder_box("HÌNH 5.2: Giao diện chi tiết sản phẩm và lựa chọn biến thể kích cỡ tại route /products/{id}")

add_para("Tại trang Giỏ hàng (Route /cart), khách điều chỉnh số lượng (hệ thống tự động kiểm tra trần tồn kho), xóa món hoặc làm trống giỏ hàng. Khi khách vãng lai đăng nhập, hệ thống tự động gộp giỏ hàng Session vào CSDL.")
add_placeholder_box("HÌNH 5.3: Giao diện giỏ hàng và cập nhật số lượng tại route /cart")

add_para("Tại trang Thanh toán (Route /checkout), khách nhập thông tin giao hàng, chọn cước vận chuyển Hỏa tốc (30.000đ) hoặc Tiêu chuẩn (15.000đ), nhập voucher giảm giá (như SAVE10, FREESHIP), chọn hình thức COD hoặc Chuyển khoản và bấm Đặt hàng. Hệ thống áp dụng idempotency token và vô hiệu hóa nút submit để chống đặt trùng đơn.")
add_placeholder_box("HÌNH 5.4: Giao diện thanh toán đơn hàng tại route /checkout")

add_para("Tại trang Lịch sử đơn hàng (Route /account/orders và /account/orders/{id}), khách xem tiến trình đơn hàng cá nhân, bấm 'Hủy đơn' nếu đơn còn ở trạng thái pending (hệ thống tự động hoàn kho), hoặc bấm 'Xác nhận đã nhận hàng' khi đơn đang shipping. Cơ chế kiểm tra quyền sở hữu ngăn chặn hoàn toàn lỗi IDOR.")
add_placeholder_box("HÌNH 5.5: Giao diện lịch sử và chi tiết đơn hàng của khách hàng tại route /account/orders")

add_sec("5.2. Phân hệ Quản trị viên (Admin Portal)")
add_para("Khu vực quản trị tại tiền tố /admin phân quyền rành mạch giữa Staff và Manager:")
add_para("Bảng điều khiển (Route /admin) cung cấp 5 thẻ thống kê: Doanh thu thực tế, Tổng đơn hàng, Đơn chờ xử lý, Tổng sản phẩm và Số lượng thành viên, đi kèm danh sách 5 đơn hàng mới nhất và đánh giá vừa gửi.")
add_placeholder_box("HÌNH 5.6: Bảng điều khiển tổng quan Admin Dashboard tại route /admin")

add_para("Quản lý Sản phẩm (Route /admin/products) cho phép tìm kiếm, thêm sản phẩm mới có upload ảnh qua Storage (/admin/products/create), chỉnh sửa giá/kho từng size (/admin/products/{id}/edit), ẩn sản phẩm (xóa mềm) và khôi phục hoạt động kinh doanh mà không ảnh hưởng đơn hàng cũ.")
add_placeholder_box("HÌNH 5.7: Giao diện quản lý sản phẩm và form thêm mới tại route /admin/products")

add_para("Quản lý Danh mục (Route /admin/categories) hiển thị số lượng sản phẩm liên kết của từng danh mục và chặn thao tác xóa nếu danh mục đang chứa sản phẩm con.")
add_placeholder_box("HÌNH 5.8: Giao diện quản lý danh mục sản phẩm tại route /admin/categories")

add_para("Quản lý Đơn hàng (Route /admin/orders và /admin/orders/{id}) cho phép lọc theo trạng thái quy trình, cập nhật tiến trình giao nhận theo ma trận ALLOWED_TRANSITIONS, xác nhận thanh toán tiền hàng và tự động hoàn trả tồn kho nếu hủy đơn.")
add_placeholder_box("HÌNH 5.9: Giao diện chi tiết đơn hàng và cập nhật trạng thái tại route /admin/orders/{id}")

add_para("Báo cáo Doanh thu (Route /admin/revenue) dành riêng cho Manager, hiển thị tổng doanh thu từ đơn delivered, giá trị đơn trung bình AOV, bảng thống kê doanh thu theo tháng và cơ cấu phương thức thanh toán.")
add_placeholder_box("HÌNH 5.10: Giao diện báo cáo thống kê doanh thu tại route /admin/revenue")

add_para("Kiểm duyệt Đánh giá (Route /admin/reviews) cho phép duyệt bật/tắt hiển thị công khai nhận xét của khách hàng hoặc xóa bỏ các bình luận spam.")
add_placeholder_box("HÌNH 5.11: Giao diện kiểm duyệt đánh giá khách hàng tại route /admin/reviews")

add_sec("5.3. Bảng đối chiếu yêu cầu kỹ thuật (Traceability Matrix)")
trace_tbl = doc.add_table(rows=17, cols=4)
trace_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(trace_tbl)

t_hdrs = ["STT", "Yêu cầu trong đề cương", "Triển khai thực tế trong mã nguồn", "Đánh giá"]
for i, h in enumerate(t_hdrs):
    c = trace_tbl.rows[0].cells[i]
    set_cell_background(c, "14221D")
    set_cell_margins(c, 50, 50, 60, 60)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

trace_data = [
    ("1", "Framework Laravel", "Laravel 11.x trên nền tảng ngôn ngữ PHP 8.3.", "ĐẠT"),
    ("2", "Cơ sở dữ liệu MySQL", "MySQL Community Server 8.4.3 với 13 bảng nghiệp vụ chuẩn hóa 3NF.", "ĐẠT"),
    ("3", "Khai báo cấu trúc bảng Migration", "4 tệp Migration khởi tạo 13 bảng, ràng buộc khóa ngoại và 4 Triggers.", "ĐẠT"),
    ("4", "Thao tác qua Eloquent ORM", "13 Models Eloquent đầy đủ quan hệ, Accessors và hỗ trợ Composite Key.", "ĐẠT"),
    ("5", "Seeder tạo dữ liệu mẫu", "ShopQuanAoSeeder nạp 841 bản ghi sạch; có file dự phòng dump SQL.", "ĐẠT"),
    ("6", "Giao diện Blade + Bootstrap", "Hệ thống Blade Templates kết hợp Bootstrap 5.2.3 và FontAwesome 6.", "ĐẠT"),
    ("7", "Tách riêng Admin và User", "Hai layout độc lập: layouts/app.blade.php và layouts/admin.blade.php.", "ĐẠT"),
    ("8", "Tối thiểu 2 vai trò người dùng", "3 vai trò Multi-Auth: Customer, Staff và Manager.", "ĐẠT"),
    ("9", "Đăng nhập / Đăng xuất Laravel Auth", "Multi-Auth Guards (web và admin), băm mật khẩu Bcrypt.", "ĐẠT"),
    ("10", "Phân quyền qua Middleware", "AdminMiddleware và ManagerOnlyMiddleware bảo vệ toàn diện /admin.", "ĐẠT"),
    ("11", "CRUD trên bảng có quan hệ", "CRUD hoàn chỉnh trên DANH_MUC và SAN_PHAM / CHI_TIET_SP, có chặn xóa.", "ĐẠT"),
    ("12", "Tìm kiếm kết hợp phân trang", "Tìm kiếm tên, lọc danh mục/giá, phân trang giữ query string withQueryString().", "ĐẠT"),
    ("13", "Validation dữ liệu nhập form", "Form Request: StoreProductRequest, CheckoutRequest... có báo lỗi tiếng Việt.", "ĐẠT"),
    ("14", "Thông báo Flash Message", "Flash session dismissible alert (success, error, info) trên cả hai layout.", "ĐẠT"),
    ("15", "Mở rộng: Upload file Storage", "Upload ảnh sản phẩm qua Storage::disk('public') kết hợp symlink.", "ĐẠT"),
    ("16", "Mở rộng: Xóa mềm / Ẩn sản phẩm", "Ẩn/khôi phục sản phẩm qua TrangThai để bảo toàn khóa ngoại đơn hàng cũ.", "ĐẠT")
]

for r_idx, (c0, c1, c2, c3) in enumerate(trace_data, start=1):
    row = trace_tbl.rows[r_idx]
    bg = "F9FBFA" if r_idx % 2 == 1 else "FFFFFF"
    for c_idx, val in enumerate([c0, c1, c2, c3]):
        c = row.cells[c_idx]
        set_cell_background(c, bg)
        set_cell_margins(c, 40, 40, 50, 50)
        p = c.paragraphs[0]
        if c_idx in [0, 3]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(val)
        if c_idx == 3:
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x00, 0x64, 0x00)
        r.font.size = Pt(9.5)

# ----------------- CHƯƠNG 6 -----------------
add_chapter("6. KIỂM THỬ VÀ ĐÁNH GIÁ CHẤT LƯỢNG SẢN PHẨM")
add_sec("6.1. Phương pháp kiểm thử tự động (Feature Tests)")
add_para("Nhóm áp dụng phương pháp kiểm thử tích hợp tự động bằng PHPUnit (php artisan test) trên cơ sở dữ liệu thử nghiệm riêng biệt smart_casual_test. Trait DatabaseTransactions được gắn vào toàn bộ 4 file Feature Tests để tự động rollback dữ liệu sau mỗi hàm test, không để lại dữ liệu rác trong database.")
add_para("Tổng cộng hệ thống có 23 kịch bản kiểm thử tự động (gồm 1 test Unit mẫu kiểm tra môi trường và 22 test Feature nghiệp vụ bao phủ toàn bộ các luồng quan trọng).")

add_sec("6.2. Kết quả kiểm thử 23 kịch bản kiểm thử tự động")
add_bullet("AdminCrudAndMiddlewareTest (5 test): ", "Kiểm tra khách vãng lai và khách thường bị chặn vào /admin; kiểm tra nhân viên staff đăng nhập thành công; kiểm tra CRUD danh mục và chặn xóa khi có sản phẩm con; kiểm tra thêm sản phẩm và xóa mềm logic.")
add_bullet("AuthenticationAndAuthorizationTest (3 test): ", "Kiểm tra đăng ký tài khoản có băm mật khẩu Bcrypt; kiểm tra đăng nhập và đăng xuất; kiểm tra bảo mật IDOR ngăn chặn khách hàng A xem trộm đơn hàng của khách hàng B (HTTP 403 Forbidden).")
add_bullet("CartAndCheckoutWorkflowTest (7 test): ", "Kiểm tra thêm/sửa/xóa giỏ hàng; kiểm tra đặt hàng nguyên tử và trừ tồn kho; kiểm tra chặn đặt hàng khi vượt quá tồn kho; kiểm tra composite key trên bảng GIO_HANG không bị ghi đè chéo giữa các mặt hàng hoặc giữa các thành viên; kiểm tra gộp giỏ hàng vãng lai khi đăng nhập; kiểm tra Database Trigger tính lại tổng tiền cho cả đơn cũ và mới.")
add_bullet("CatalogAndProductTest (7 test): ", "Kiểm tra quan hệ Eloquent; kiểm tra tải trang chủ; kiểm tra tìm kiếm theo từ khóa; kiểm tra lọc theo danh mục; kiểm tra phân trang giữ query string; kiểm tra chi tiết sản phẩm và gửi đánh giá nhận xét.")

add_para("Log thực thi thực tế khi chạy lệnh php artisan test trong Terminal:")
p_code = doc.add_paragraph()
p_code.paragraph_format.left_indent = Inches(0.2)
p_code.paragraph_format.space_before = Pt(4)
p_code.paragraph_format.space_after = Pt(6)
r_c = p_code.add_run("""   PASS  Tests\\Unit\\ExampleTest
  ✓ that true is true                                                                         0.01s

   PASS  Tests\\Feature\\AdminCrudAndMiddlewareTest
  ✓ guest cannot access admin dashboard                                                      0.20s
  ✓ customer cannot access admin dashboard                                                   0.03s
  ✓ staff can login and access admin dashboard                                               0.04s
  ✓ admin category crud workflow                                                             0.06s
  ✓ admin product crud and soft delete                                                       0.06s

   PASS  Tests\\Feature\\AuthenticationAndAuthorizationTest
  ✓ customer registration with password hashing                                              0.03s
  ✓ customer login and logout                                                                0.23s
  ✓ user cannot access other users order idor protection                                     0.05s

   PASS  Tests\\Feature\\CartAndCheckoutWorkflowTest
  ✓ cart add update remove operations                                                        0.04s
  ✓ checkout atomic workflow and stock deduction                                             0.03s
  ✓ cannot checkout more than available stock                                                0.03s
  ✓ member composite cart with three items update and delete                                 0.04s
  ✓ two members with same variant do not interfere                                           0.04s
  ✓ guest cart and member cart merge cleanly                                                 0.05s
  ✓ trigger update order items recalculates both old and new orders                          0.05s

   PASS  Tests\\Feature\\CatalogAndProductTest
  ✓ database relationships between category and product                                      0.03s
  ✓ storefront home page loads with products                                                 0.05s
  ✓ search products by keyword                                                               0.04s
  ✓ filter products by category                                                              0.06s
  ✓ pagination and query string retention                                                    0.06s
  ✓ product detail page loads                                                                0.04s
  ✓ submit product review                                                                    0.04s

  Tests:    23 passed (118 assertions)
  Duration: 1.42s""")
r_c.font.name = 'Consolas'
r_c.font.size = Pt(9.5)
r_c.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

add_para("Phân tích kết quả: Toàn bộ 23/23 kịch bản kiểm thử đều đạt kết quả PASS với 118 assertions trong 1.42 giây. Kết quả này chứng minh hệ thống phân quyền chặn đúng đối tượng, các giao dịch tài chính và trừ kho được thực hiện nguyên tử an toàn, và các triggers tự động hóa cập nhật số liệu chính xác.")

# ----------------- CHƯƠNG 7 -----------------
add_chapter("7. PHÂN CÔNG CÔNG VIỆC VÀ ĐÁNH GIÁ ĐÓNG GÓP NHÓM")
add_sec("7.1. Bảng phân công nhiệm vụ chi tiết")
assign_tbl = doc.add_table(rows=4, cols=5)
assign_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(assign_tbl)

a_hdrs = ["STT", "Họ và tên", "MSSV - Lớp", "Nhiệm vụ chính đảm nhiệm", "Mức độ hoàn thành"]
for i, h in enumerate(a_hdrs):
    c = assign_tbl.rows[0].cells[i]
    set_cell_background(c, "14221D")
    set_cell_margins(c, 50, 50, 60, 60)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

assign_data = [
    ("1", "Nguyễn Huy Hoàng", "2001230266\n14DHTH02", "Khởi tạo project Laravel 11; Viết Migration 13 bảng & 4 Triggers; Xây dựng 13 Models Eloquent; Triển khai CheckoutService (lockForUpdate) & CartService; Lập trình Admin Controllers; Viết trọn bộ 23 kịch bản kiểm thử tự động.", "100%"),
    ("2", "Nguyễn Việt Vương", "2001231062\n14DHTH02", "Chuyển đổi giao diện HTML sang Blade Templates; Xây dựng Layout Storefront & Admin; Thiết kế trang Catalog, Chi tiết sản phẩm, Giỏ hàng, Thanh toán; Đảm bảo tương thích di động Bootstrap 5; Xử lý JavaScript phía client.", "100%"),
    ("3", "Nguyễn Ngọc Tường Vi", "2001231032\n14DHTH12", "Thu thập và chuẩn hóa dữ liệu Seeder (841 bản ghi vào 13 bảng); Lập trình module Khách hàng đăng ký, đăng nhập; Soạn thảo báo cáo đồ án kỹ thuật 10 chương; Chuẩn bị tài liệu hướng dẫn README và slide thuyết trình.", "100%")
]

for r_idx, (c0, c1, c2, c3, c4) in enumerate(assign_data, start=1):
    row = assign_tbl.rows[r_idx]
    bg = "F9FBFA" if r_idx % 2 == 1 else "FFFFFF"
    for c_idx, val in enumerate([c0, c1, c2, c3, c4]):
        c = row.cells[c_idx]
        set_cell_background(c, bg)
        set_cell_margins(c, 40, 40, 50, 50)
        p = c.paragraphs[0]
        if c_idx in [0, 4]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(val)
        if c_idx == 4:
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x00, 0x64, 0x00)
        r.font.size = Pt(9.5)

add_sec("7.2. Tỷ lệ hoàn thành và phương thức phối hợp nhóm")
add_para("Về mức độ đóng góp: Cả ba thành viên trong nhóm đều hoàn thành 100% khối lượng công việc được phân công, có tinh thần trách nhiệm cao và chủ động phối hợp giải quyết các vướng mắc phát sinh trong quá trình xây dựng ứng dụng.")
add_para("Về phương thức phối hợp qua Git: Nhóm sử dụng GitHub làm nền tảng quản lý mã nguồn tập trung. Nhánh main là nhánh mã nguồn ổn định, chỉ hợp nhất khi các tính năng đã vượt qua bộ 23 kịch bản kiểm thử tự động. Mỗi thành viên khi phát triển module mới đều tạo nhánh riêng từ main, tự rebase cục bộ trước khi tạo Pull Request để hạn chế tối đa xung đột mã nguồn. Trước khi lập trình CSDL, nhóm họp thống nhất danh sách tên bảng và tên cột chuẩn, nhờ đó mã nguồn giao diện Blade và controller backend luôn ăn khớp hoàn toàn.")

# ----------------- CHƯƠNG 8 -----------------
add_chapter("8. HẠN CHẾ VÀ HƯỚNG PHÁT TRIỂN")
add_sec("8.1. Các hạn chế hiện tại")
add_bullet("1. Chưa tích hợp cổng thanh toán trực tuyến tự động: ", "Mới hỗ trợ COD và chuyển khoản ngân hàng thủ công (nhân viên phải tự kiểm tra tài khoản rồi bấm xác nhận đã thanh toán trên trang quản trị); chưa kết nối với VNPay hay MoMo qua Webhook IPN để cập nhật trạng thái tự động.")
add_bullet("2. Chưa có dịch vụ gửi email tự động: ", "Chưa cấu hình Laravel Mail qua SMTP/Mailpit để gửi email xác nhận hóa đơn cho khách sau khi đặt hàng và chưa có luồng đặt lại mật khẩu qua email xác thực.")
add_bullet("3. Chưa áp dụng bộ nhớ đệm (Caching): ", "Toàn bộ truy vấn trang chủ và catalog đang đọc trực tiếp từ MySQL qua Eloquent. Khi lượng truy cập tăng cao, việc quét bảng kèm nhiều điều kiện lọc và tính sao trung bình có thể làm tăng thời gian phản hồi.")
add_bullet("4. Bảng mã giảm giá thiếu thời hạn hiệu lực: ", "Bảng MA_GIAM_GIA chưa có cột ngày bắt đầu và kết thúc; việc hết hạn hoàn toàn phụ thuộc vào việc chuyển trạng thái thủ công hoặc khi số lượt dùng giảm về 0.")

add_sec("8.2. Kế hoạch nâng cấp và phát triển mở rộng")
add_para("Lộ trình 4 giai đoạn nâng cấp hệ thống:")
add_bullet("Giai đoạn 1 (Ưu tiên cao nhất): ", "Tích hợp cổng thanh toán VNPay Sandbox với Webhook IPN tự động nhận diện thanh toán không tiền mặt.")
add_bullet("Giai đoạn 2 (Trải nghiệm người dùng): ", "Cấu hình hàng đợi Laravel Queue và gửi email hóa đơn PDF tự động qua Laravel Mail.")
add_bullet("Giai đoạn 3 (Tối ưu hiệu năng): ", "Tích hợp Redis Cache để lưu tạm danh mục và sản phẩm thịnh hành vào RAM, giảm tải câu lệnh SELECT xuống MySQL.")
add_bullet("Giai đoạn 4 (Mở rộng nền tảng): ", "Bổ sung thời hạn cho mã giảm giá và xây dựng RESTful API cho ứng dụng di động Flutter/React Native.")

# ----------------- CHƯƠNG 9 -----------------
add_chapter("9. KẾT LUẬN")
add_para("Dự án \"SMART CASUAL — Website Bán hàng Thời trang bằng Laravel Framework và MySQL\" đã hoàn thành đầy đủ các mục tiêu đề ra trong học phần Lập trình mã nguồn mở tại Trường Đại học Công Thương TP. Hồ Chí Minh (HUIT). Hệ thống được xây dựng hoàn chỉnh trên nền tảng Laravel 11.x kết hợp cơ sở dữ liệu quan hệ MySQL 8.4, đáp ứng trọn vẹn chu trình mua sắm thời trang trực tuyến với 13 bảng nghiệp vụ chuẩn hóa, từ quản lý danh mục, biến thể kích cỡ sản phẩm, giỏ hàng, đặt hàng đến bảng điều khiển quản trị phân quyền cho nhân viên và quản lý cửa hàng.")
add_para("Trong quá trình thực hiện đồ án, nhóm đã giải quyết được các khó khăn kỹ thuật then chốt: xử lý bài toán tranh chấp dữ liệu (Race Condition) khi đặt hàng bằng cơ chế khóa dòng dữ liệu bi quan lockForUpdate() kết hợp Database Transaction tại lớp dịch vụ CheckoutService, đồng thời khắc phục hạn chế của Eloquent ORM đối với khóa chính phức hợp hai cột trên bảng GIO_HANG bằng cách ghi đè phương thức setKeysForSaveQuery.")
add_para("Thông qua đồ án này, các thành viên trong nhóm đã nắm vững phương pháp tổ chức mã nguồn theo mô hình MVC kết hợp Service Layer, hiểu rõ cách thức phân quyền người dùng qua Multi-Auth Guards và Middleware tùy chỉnh, đồng thời làm quen với việc phối hợp kỹ thuật giữa tầng ứng dụng và tầng cơ sở dữ liệu qua Database Triggers. Bộ 23 kịch bản kiểm thử tự động được xây dựng cũng giúp nhóm rèn luyện kỹ năng kiểm thử phần mềm có phương pháp, đảm bảo các tính năng vận hành ổn định.")

# ----------------- CHƯƠNG 10 -----------------
add_chapter("10. TÀI LIỆU THAM KHẢO")
refs = [
    "[1] Khoa Công nghệ Thông tin — Trường Đại học Công Thương TP. HCM, Đề cương chi tiết học phần Lập trình mã nguồn mở, Bộ môn Kỹ thuật Phần mềm, 2026.",
    "[2] T. Otwell, Laravel Documentation: The PHP Framework for Web Artisans (Version 11.x), Laravel LLC, 2026. [Trực tuyến]. Khả dụng: https://laravel.com/docs/11.x",
    "[3] The PHP Group, PHP 8.3 Manual and Language Reference, 2026. [Trực tuyến]. Khả dụng: https://www.php.net/manual/en/",
    "[4] Oracle Corporation, MySQL 8.4 Reference Manual, Oracle Corporation, 2026. [Trực tuyến]. Khả dụng: https://dev.mysql.com/doc/refman/8.4/en/",
    "[5] M. Otto và J. Thornton, Bootstrap 5.2 Documentation: The most popular HTML, CSS, and JS library in the world, Bootstrap Core Team, 2026. [Trực tuyến]. Khả dụng: https://getbootstrap.com/docs/5.2/",
    "[6] S. Bergmann, PHPUnit 10 Manual: The PHP Testing Framework, 2026. [Trực tuyến]. Khả dụng: https://phpunit.de/documentation.html"
]

for r_text in refs:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(r_text)
    r.font.size = Pt(11)

# Save document
output_path = Path(r"D:\PHP_NOSQL\SMART_CASUAL_LARAVEL\docs\BAO_CAO_DO_AN_MON_HOC_LARAVEL_HOAN_CHINH.docx")
doc.save(str(output_path))
print(f"Successfully compiled: {output_path}")
