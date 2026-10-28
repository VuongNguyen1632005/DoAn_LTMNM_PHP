import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pathlib import Path

def test_page_number():
    doc = docx.Document()
    sec = doc.sections[0]
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Trang ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    
    # Page field
    run_page = p.add_run()
    run_page.font.name = "Times New Roman"
    run_page.font.size = Pt(10)
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run_page._r.append(fld1)
    run_page._r.append(instr)
    run_page._r.append(fld2)
    run_page._r.append(fld3)

    run_of = p.add_run(" / ")
    run_of.font.name = "Times New Roman"
    run_of.font.size = Pt(10)

    # Numpages field
    run_numpages = p.add_run()
    run_numpages.font.name = "Times New Roman"
    run_numpages.font.size = Pt(10)
    fld1_np = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr_np = parse_xml(r'<w:instrText %s xml:space="preserve"> NUMPAGES </w:instrText>' % nsdecls('w'))
    fld2_np = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3_np = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run_numpages._r.append(fld1_np)
    run_numpages._r.append(instr_np)
    run_numpages._r.append(fld2_np)
    run_numpages._r.append(fld3_np)

    doc.add_paragraph("Trang 1 thử nghiệm.")
    doc.add_page_break()
    doc.add_paragraph("Trang 2 thử nghiệm.")
    
    out_path = Path("D:/PHP_NOSQL/SMART_CASUAL_LARAVEL/docs/test_page_num.docx")
    doc.save(str(out_path))
    print("Page number test succeeded:", out_path)

if __name__ == "__main__":
    test_page_number()
