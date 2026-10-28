import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

p = r"C:\Users\VuongNguyen\.gemini\antigravity\brain\6f1d1daf-5c59-4014-bff0-7e3389b6bf51\.user_uploaded\media_1793017301535_b532377e.docx"
doc = docx.Document(p)

for para in doc.paragraphs:
    if para.text.strip():
        print(para.text)

for table in doc.tables:
    for row in table.rows:
        print(" | ".join([cell.text.replace("\n", " ").strip() for cell in row.cells]))
