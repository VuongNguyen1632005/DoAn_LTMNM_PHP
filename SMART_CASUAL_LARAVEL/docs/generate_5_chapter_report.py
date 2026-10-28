import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from pathlib import Path

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="B0C4DE"):
    tblPr = table._element.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

doc = docx.Document()

# Page setup: A4, Margins: Top 2cm, Bottom 2cm, Left 2.5cm, Right 2cm
for section in doc.sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.79)
    section.bottom_margin = Inches(0.79)
    section.left_margin = Inches(0.98)
    section.right_margin = Inches(0.79)

normal_style = doc.styles['Normal']
normal_style.font.name = 'Times New Roman'
normal_style.font.size = Pt(13)
normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
normal_style.paragraph_format.line_spacing = 1.3
normal_style.paragraph_format.space_after = Pt(4)

# ----------------- COVER PAGE -----------------
header_tbl = doc.add_table(rows=1, cols=2)
header_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
header_tbl.autofit = False

c0 = header_tbl.rows[0].cells[0]
c0.width = Inches(3.2)
p0 = c0.paragraphs[0]
p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p0.add_run("TRƯỜNG ĐH CÔNG THƯƠNG TP. HCM\n")
r.font.bold = True
r.font.size = Pt(11)
r2 = p0.add_run("KHOA CÔNG NGHỆ THÔNG TIN\nBỘ MÔN KỸ THUẬT PHẦN MỀM\n❖❖❖")
r2.font.size = Pt(10.5)

c1 = header_tbl.rows[0].cells[1]
c1.width = Inches(3.5)
p1 = c1.paragraphs[0]
p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p1.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
r.font.bold = True
r.font.size = Pt(11)
r2 = p1.add_run("Độc lập – Tự do – Hạnh phúc\n")
r2.font.bold = True
r2.font.size = Pt(10.5)
r3 = p1.add_run("— — — — — — — — —")
r3.font.size = Pt(10)

doc.add_paragraph().paragraph_format.space_after = Pt(28)

title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_t1 = title_p.add_run("BÁO CÁO ĐỒ ÁN CUỐI MÔN HỌC\n")
r_t1.font.bold = True
r_t1.font.size = Pt(15)
r_t1.font.color.rgb = RGBColor(0x14, 0x22, 0x1D)

r_t2 = title_p.add_run("HỌC PHẦN: LẬP TRÌNH MÃ NGUỒN MỞ\n\n")
r_t2.font.bold = True
r_t2.font.size = Pt(13)

r_t3 = title_p.add_run("ĐỀ TÀI:\nXÂY DỰNG WEBSITE BÁN HÀNG THỜI TRANG \"SMART CASUAL\"\nBẰNG LARAVEL FRAMEWORK VÀ CƠ SỞ DỮ LIỆU MYSQL\n")
r_t3.font.bold = True
r_t3.font.size = Pt(16)
r_t3.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

doc.add_paragraph().paragraph_format.space_after = Pt(36)

# Information table
info_tbl = doc.add_table(rows=6, cols=2)
info_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(info_tbl, "14221D")

info_data = [
    ("Sinh viên 1 (Trưởng nhóm):", "Nguyễn Huy Hoàng — MSSV: 2001230266 — Lớp: 14DHTH02"),
    ("Sinh viên 2 (Thành viên):", "Nguyễn Việt Vương — MSSV: 2001231062 — Lớp: 14DHTH02"),
    ("Sinh viên 3 (Thành viên):", "Nguyễn Ngọc Tường Vi — MSSV: 2001231032 — Lớp: 14DHTH12"),
    ("Giảng viên hướng dẫn:", "Thầy Nguyễn Minh Hiến"),
    ("Ngành đào tạo:", "Công nghệ Thông tin — Chuyên ngành Kỹ thuật Phần mềm"),
    ("Thời gian thực hiện:", "Học kỳ 1 — Năm học 2026 – 2027")
]

for idx, (label, val) in enumerate(info_data):
    row = info_tbl.rows[idx]
    cell0 = row.cells[0]
    cell0.width = Inches(2.2)
    set_cell_background(cell0, "F4F6F8")
    set_cell_margins(cell0, 70, 70, 100, 100)
    p = cell0.paragraphs[0]
    r = p.add_run(label)
    r.font.bold = True
    r.font.size = Pt(11)
    
    cell1 = row.cells[1]
    cell1.width = Inches(4.5)
    set_cell_margins(cell1, 70, 70, 100, 100)
    p = cell1.paragraphs[0]
    r = p.add_run(val)
    r.font.size = Pt(11)

doc.add_page_break()

# ----------------- LỜI CẢM ƠN -----------------
h_ack = doc.add_paragraph()
h_ack.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h_ack.add_run("LỜI CẢM ƠN\n")
r.font.bold = True
r.font.size = Pt(15)

p_ack = doc.add_paragraph()
p_ack.paragraph_format.line_spacing = 1.3
p_ack.paragraph_format.space_after = Pt(8)
p_ack.add_run("Lời đầu tiên, nhóm sinh viên chúng em xin gửi lời cảm ơn sâu sắc đến Thầy Nguyễn Minh Hiến, giảng viên phụ trách học phần Lập trình mã nguồn mở, Bộ môn Kỹ thuật Phần mềm, Khoa Công nghệ Thông tin, Trường Đại học Công Thương TP. Hồ Chí Minh. Trong suốt thời gian học tập và thực hiện đề tài, Thầy đã truyền đạt những kiến thức nền tảng về framework Laravel, hướng dẫn cấu trúc mã nguồn chuẩn mực và đưa ra các định hướng thiết thực giúp nhóm giải quyết các vướng mắc về cơ sở dữ liệu quan hệ, kỹ thuật khóa dòng giao dịch cũng như cơ chế phân quyền người dùng.")

p_ack2 = doc.add_paragraph()
p_ack2.paragraph_format.line_spacing = 1.3
p_ack2.paragraph_format.space_after = Pt(8)
p_ack2.add_run("Nhóm cũng xin chân thành cảm ơn các Thầy Cô trong Khoa Công nghệ Thông tin đã tận tâm giảng dạy, trang bị cho chúng em nền tảng tư duy lập trình vững chắc qua các môn học chuyên ngành. Dù các thành viên đã nỗ lực hoàn thành dự án, sản phẩm và báo cáo chắc chắn khó tránh khỏi những điểm thiếu sót. Nhóm kính mong nhận được những góp ý quý báu của Thầy Cô để hệ thống ngày càng hoàn thiện hơn.")

doc.add_paragraph().paragraph_format.space_after = Pt(16)

# ----------------- NHẬN XÉT CỦA GIẢNG VIÊN -----------------
h_rev = doc.add_paragraph()
h_rev.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h_rev.add_run("NHẬN XÉT CỦA GIẢNG VIÊN HƯỚNG DẪN\n")
r.font.bold = True
r.font.size = Pt(14)

for _ in range(8):
    p_dot = doc.add_paragraph()
    p_dot.paragraph_format.space_after = Pt(3)
    p_dot.add_run("................................................................................................................................................................................................")

p_score = doc.add_paragraph()
p_score.paragraph_format.space_before = Pt(8)
p_score.add_run("Điểm đánh giá (bằng số): ....................       Điểm đánh giá (bằng chữ): ....................................")

p_sign = doc.add_paragraph()
p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_sign.paragraph_format.space_before = Pt(12)
r = p_sign.add_run("Tp. Hồ Chí Minh, ngày ..... tháng ..... năm 2026\n")
r.font.italic = True
r2 = p_sign.add_run("Giảng viên đánh giá\n")
r2.font.bold = True
r3 = p_sign.add_run("(Ký và ghi rõ họ tên)\n\n\n\n\nThầy Nguyễn Minh Hiến")

doc.add_page_break()

# ----------------- MỤC LỤC 5 CHƯƠNG -----------------
h_toc = doc.add_paragraph()
h_toc.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h_toc.add_run("MỤC LỤC BÁO CÁO\n")
r.font.bold = True
r.font.size = Pt(15)

toc_items = [
    ("CHƯƠNG 1. GIỚI THIỆU ĐỀ TÀI", True),
    ("   1.1. Lý do chọn đề tài", False),
    ("   1.2. Mục tiêu đề tài", False),
    ("   1.3. Phạm vi thực hiện", False),
    ("CHƯƠNG 2. PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG", True),
    ("   2.1. Phân tích đối tượng và vai trò người dùng", False),
    ("   2.2. Sơ đồ UML Use Case tổng quan", False),
    ("   2.3. Sơ đồ và thiết kế Cơ sở dữ liệu (ERD)", False),
    ("   2.4. Danh mục 13 bảng nghiệp vụ chi tiết", False),
    ("   2.5. Thiết kế Triggers và tính toán tự động", False),
    ("CHƯƠNG 3. CÔNG NGHỆ, KIẾN TRÚC VÀ TRIỂN KHAI", True),
    ("   3.1. Kiến trúc tổng thể hệ thống", False),
    ("   3.2. Danh mục công nghệ và thư viện sử dụng", False),
    ("   3.3. Cơ chế xác thực và phân quyền (Multi-Auth & Middleware)", False),
    ("   3.4. Mô hình ORM Eloquent và bảo vệ toàn vẹn dữ liệu", False),
    ("   3.5. Yêu cầu môi trường hệ thống", False),
    ("   3.6. Các bước cài đặt chi tiết từ mã nguồn", False),
    ("   3.7. Danh sách tài khoản thử nghiệm hệ thống", False),
    ("CHƯƠNG 4. KẾT QUẢ VÀ KIỂM THỬ HỆ THỐNG", True),
    ("   4.1. Phân hệ Khách hàng (Customer Storefront)", False),
    ("   4.2. Phân hệ Quản trị viên (Admin Portal)", False),
    ("   4.3. Bảng đối chiếu yêu cầu kỹ thuật (Traceability Matrix)", False),
    ("   4.4. Phương pháp kiểm thử tự động (Feature Tests)", False),
    ("   4.5. Kết quả kiểm thử tự động", False),
    ("CHƯƠNG 5. ĐÁNH GIÁ NHÓM, HẠN CHẾ VÀ HƯỚNG PHÁT TRIỂN", True),
    ("   5.1. Bảng phân công nhiệm vụ chi tiết", False),
    ("   5.2. Tỷ lệ hoàn thành công việc", False),
    ("   5.3. Các hạn chế hiện tại", False),
    ("   5.4. Kế hoạch nâng cấp và phát triển mở rộng", False),
    ("KẾT LUẬN", True),
    ("TÀI LIỆU THAM KHẢO", True),
]

for title, is_bold in toc_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(title)
    r.font.bold = is_bold
    r.font.size = Pt(11.5 if is_bold else 11)

doc.add_page_break()

# ----------------- DANH MỤC -----------------
h_abbr = doc.add_paragraph()
h_abbr.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h_abbr.add_run("DANH MỤC TỪ VIẾT TẮT\n")
r.font.bold = True
r.font.size = Pt(13)

abbr_list = [
    ("3NF", "Third Normal Form — Dạng chuẩn 3 trong thiết kế cơ sở dữ liệu quan hệ"),
    ("AOV", "Average Order Value — Giá trị đơn hàng trung bình của các giao dịch"),
    ("COD", "Cash On Delivery — Phương thức thanh toán bằng tiền mặt khi nhận hàng"),
    ("CRUD", "Create, Read, Update, Delete — Bốn thao tác dữ liệu cơ bản"),
    ("CSRF", "Cross-Site Request Forgery — Tấn công giả mạo yêu cầu từ phía người dùng tin cậy"),
    ("ERD", "Entity Relationship Diagram — Sơ đồ quan hệ thực thể mô tả cấu trúc CSDL"),
    ("IDOR", "Insecure Direct Object References — Lỗ hổng tham chiếu trực tiếp đối tượng"),
    ("MVC", "Model - View - Controller — Mô hình kiến trúc phân tách Dữ liệu - Giao diện - Điều hướng"),
    ("ORM", "Object-Relational Mapping — Kỹ thuật ánh xạ dữ liệu quan hệ sang đối tượng"),
    ("PDO", "PHP Data Objects — Giao diện truy cập cơ sở dữ liệu chuẩn trong PHP"),
    ("RBAC", "Role-Based Access Control — Kiểm soát truy cập dựa trên vai trò"),
    ("SKU", "Stock Keeping Unit — Đơn vị phân loại hàng hóa lưu kho theo biến thể"),
    ("XSS", "Cross-Site Scripting — Hình thức tấn công chèn mã kịch bản độc hại vào trang web")
]

for a, exp in abbr_list:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"• {a}: ")
    r.font.bold = True
    p.add_run(exp)

doc.add_paragraph().paragraph_format.space_after = Pt(10)

h_tbl_list = doc.add_paragraph()
h_tbl_list.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h_tbl_list.add_run("DANH MỤC BẢNG BIỂU & HÌNH ẢNH\n")
r.font.bold = True
r.font.size = Pt(13)

fig_tbl_items = [
    "Bảng 3.1: Danh sách tài khoản thử nghiệm hệ thống Smart Casual",
    "Bảng 4.1: Bảng đối chiếu yêu cầu kỹ thuật trong đề cương môn học (Traceability Matrix)",
    "Bảng 5.1: Bảng phân công nhiệm vụ chi tiết của các thành viên trong nhóm",
    "Bảng 5.2: Bảng đánh giá mức độ hoàn thành nhiệm vụ của các thành viên",
    "Hình 2.1: Sơ đồ Use Case tổng quan thể hiện tương tác của từng vai trò trong hệ thống Smart Casual",
    "Hình 2.2: Sơ đồ quan hệ thực thể (ERD) của hệ thống thương mại điện tử Smart Casual",
    "Hình 2.3: Sơ đồ phối hợp trách nhiệm giữa Service Layer của ứng dụng và Database Triggers trong MySQL",
    "Hình 3.1: Sơ đồ luồng xử lý HTTP request qua các tầng kiến trúc của hệ thống Smart Casual",
    "Hình 4.1: Giao diện trang chủ và bộ lọc catalog sản phẩm tại route /",
    "Hình 4.2: Giao diện chi tiết sản phẩm và lựa chọn biến thể kích cỡ tại route /products/{id}",
    "Hình 4.3: Giao diện giỏ hàng và cập nhật số lượng tại route /cart",
    "Hình 4.4: Giao diện thanh toán đơn hàng tại route /checkout",
    "Hình 4.5: Giao diện lịch sử và chi tiết đơn hàng của khách hàng tại route /account/orders",
    "Hình 4.6: Bảng điều khiển tổng quan Admin Dashboard tại route /admin",
    "Hình 4.7: Giao diện quản lý sản phẩm và form thêm mới tại route /admin/products",
    "Hình 4.8: Giao diện quản lý danh mục sản phẩm tại route /admin/categories",
    "Hình 4.9: Giao diện chi tiết đơn hàng và cập nhật trạng thái tại route /admin/orders/{id}",
    "Hình 4.10: Giao diện báo cáo thống kê doanh thu tại route /admin/revenue",
    "Hình 4.11: Giao diện kiểm duyệt đánh giá khách hàng tại route /admin/reviews",
    "Hình 5.1: Lộ trình 4 giai đoạn nâng cấp và mở rộng hệ thống Smart Casual"
]

for item in fig_tbl_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.add_run(f"• {item}")

doc.add_page_break()

# ----------------- HELPERS -----------------
def add_chapter(num_title):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(num_title)
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x14, 0x22, 0x1D)
    return h

def add_sec(title):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(title)
    r.font.bold = True
    r.font.size = Pt(12.5)
    r.font.color.rgb = RGBColor(0x24, 0x3B, 0x32)
    return h

def add_para(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run(text)
    r.font.size = Pt(12.5)
    return p

def add_caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

def add_tbl_caption(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.bold = True
    r.font.size = Pt(11.5)
    r.font.color.rgb = RGBColor(0x14, 0x22, 0x1D)

def add_box(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"[{text}]")
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

# ----------------- CHƯƠNG 1 -----------------
add_chapter("CHƯƠNG 1. GIỚI THIỆU ĐỀ TÀI")
add_para("Chương này trình bày bối cảnh thực tế dẫn đến việc lựa chọn đề tài xây dựng website bán hàng thời trang Smart Casual. Tiếp đó, chương xác định các mục tiêu cụ thể cần đạt được theo chuẩn đầu ra của học phần và phạm vi chức năng mà hệ thống hướng tới phục vụ các nhóm người dùng khác nhau.")

add_sec("1.1. Lý do chọn đề tài")
add_para("Các cửa hàng kinh doanh thời trang bán lẻ quy mô vừa và nhỏ thường gặp khó khăn trong việc kiểm soát số lượng tồn kho theo từng kích cỡ sản phẩm. Khi nhập trang phục hoặc giày dép về kinh doanh, mỗi mẫu mã luôn đi kèm nhiều kích cỡ như size S, M, L, XL đối với áo quần hoặc 38 đến 43 đối với giày dép. Mỗi kích cỡ có lượng hàng trong kho riêng biệt và đôi khi có mức giá bán khác nhau giữa các size phổ thông và size ngoại cỡ. Nhiều cửa hàng hiện vẫn ghi chép hàng hóa qua sổ tay hoặc tiếp nhận đơn hàng phân tán qua các tin nhắn trên mạng xã hội, dẫn đến tình trạng người bán không nắm bắt được số lượng tồn thực tế của từng size tại thời điểm khách hỏi mua. Hậu quả là nhân viên xác nhận đơn cho những kích cỡ đã hết hàng từ trước, buộc phải liên hệ lại để hủy đơn hoặc thuyết phục khách đổi mẫu, gây phiền toái và làm giảm sự tin tưởng của người tiêu dùng.")
add_para("Vấn đề tiếp theo nằm ở quy trình luân chuyển đơn hàng và đối soát tài chính. Một đơn hàng sau khi khách đặt mua phải trải qua nhiều bước xử lý: nhân viên kho kiểm tra sản phẩm, đóng gói kiện hàng, bàn giao cho đơn vị vận chuyển, thu tiền mặt khi giao hàng (COD) và ghi nhận thanh toán vào hệ thống. Nếu không có một ứng dụng phân định trạng thái đơn rõ ràng, nhân viên phụ trách bán hàng và nhân sự kho thường xuyên phải nhắn tin hay gọi điện nội bộ để kiểm tra tiến độ giao hàng. Sự thiếu đồng bộ này dễ dẫn đến nhầm lẫn trạng thái, chậm trễ thời gian phát hàng và gây khó khăn khi đối chiếu doanh thu thực thu với các đối tác vận chuyển vào cuối tháng.")
add_para("Bên cạnh đó, các cửa hàng nhỏ thường không đủ ngân sách để trang bị các giải pháp phần mềm doanh nghiệp cồng kềnh với chi phí bản quyền định kỳ cao. Họ cần một hệ thống bán hàng trực tuyến vừa đủ gọn nhẹ, giao diện rõ ràng cho khách mua hàng trên điện thoại, đồng thời cung cấp khu vực quản trị đơn giản để nhân viên kho kiểm tra tồn kho theo size và cập nhật tiến trình đơn hàng. Nền tảng mã nguồn mở Laravel kết hợp với hệ quản trị cơ sở dữ liệu quan hệ MySQL là lựa chọn phù hợp để xây dựng ứng dụng này nhờ khả năng tùy biến cao, cấu trúc lập trình rõ ràng và chi phí triển khai thấp.")
add_para("Xuất phát từ các bài toán nghiệp vụ thực tế trên, nhóm chọn đề tài \"Xây dựng Website Bán hàng Thời trang Smart Casual bằng Laravel Framework và MySQL\". Đề tài tập trung giải quyết bài toán kinh doanh trang phục thường ngày (smart casual), tổ chức dữ liệu biến thể theo kích cỡ, chuẩn hóa luồng đặt hàng và tính toán tồn kho. Thông qua việc phát triển hệ thống, các thành viên trong nhóm có điều kiện vận dụng kiến thức của học phần Lập trình mã nguồn mở, từ tư duy hướng đối tượng trong PHP, mô hình MVC, truy vấn dữ liệu quan hệ qua Eloquent ORM đến kỹ thuật kiểm soát truy cập bằng Middleware vào một bài toán thực tế.")

add_sec("1.2. Mục tiêu đề tài")
add_para("Dự án hướng đến việc ứng dụng công nghệ mã nguồn mở để giải quyết bài toán kinh doanh cho một cửa hàng thời trang bán lẻ. Các mục tiêu của đồ án được phân định thành hai nhóm chính, bao gồm việc đáp ứng chuẩn đầu ra theo đề cương môn học và việc hoàn thiện các tiêu chuẩn kỹ thuật phần mềm cho hệ thống.")
add_para("Về năng lực lập trình và xây dựng giao diện, nhóm vận dụng ngôn ngữ PHP 8.3 cùng framework Laravel 11.x và cơ sở dữ liệu MySQL để hiện thực hóa các chức năng của một website thương mại điện tử (CLO1). Giao diện người dùng được xây dựng bằng Blade Template kết hợp Bootstrap 5.2.3, có khả năng co giãn trên màn hình máy tính lẫn điện thoại, cung cấp thông báo trạng thái cụ thể sau mỗi thao tác thêm hàng, sửa giỏ hoặc cập nhật đơn (CLO2).")
add_para("Về phương thức phối hợp nhóm, các thành viên chia việc theo từng mảng chuyên môn gồm backend, giao diện và kiểm thử tài liệu. Mã nguồn dự án được quản lý tập trung trên GitHub, trong đó các tính năng mới đều được phát triển trên nhánh riêng, trải qua quá trình kiểm tra mã và thử nghiệm trước khi hợp nhất vào nhánh chính (CLO3).")
add_para("Về kết quả sản phẩm và tinh thần tự học, nhóm đặt mục tiêu xây dựng một ứng dụng có thể vận hành thực tế cho ngành hàng trang phục với các luồng duyệt hàng, đặt hàng, xử lý giao nhận và thống kê doanh thu (CLO4). Trong suốt quá trình thực hiện, nhóm tự tìm hiểu các cơ chế xử lý tranh chấp tồn kho, kiểm soát quyền truy cập và tuân thủ các quy định về trích dẫn tài liệu tham khảo (CLO5).")
add_para("Đối với cơ sở dữ liệu, mục tiêu của nhóm là chuẩn hóa 13 bảng theo dạng chuẩn 3NF và thiết lập 4 Database Triggers. Đối với một cửa hàng thời trang có nhiều mặt hàng và thường xuyên thay đổi số lượng mua, việc để cơ sở dữ liệu tự động tính lại tổng tiền qua trigger giúp số liệu trên đơn hàng luôn khớp đúng với các dòng chi tiết mà không phụ thuộc vào code ứng dụng.")
add_para("Về kiến trúc phần mềm, nhóm tổ chức mã nguồn theo mô hình MVC kết hợp lớp dịch vụ Service Layer của Laravel 11.x. Tầng Controller chỉ làm nhiệm vụ điều phối request, tầng Form Request kiểm tra tính hợp lệ của dữ liệu đầu vào, trong khi tầng Service Layer đảm nhiệm logic tính toán chiết khấu và trừ tồn kho. Cách phân chia này giúp mã nguồn không bị dồn quá nhiều xử lý vào Controller, thuận tiện cho việc kiểm thử và bảo trì sau này.")
add_para("Về mặt an toàn hệ thống, nhóm đặt mục tiêu phân tách hai guard xác thực độc lập cho khách mua hàng và nhân sự quản lý, kết hợp hệ thống Middleware kiểm tra vai trò. Do nhân viên cửa hàng và người mua có thông tin tài khoản cùng phạm vi thao tác khác biệt, việc tách guard giúp phiên đăng nhập không bị nhầm lẫn. Bên cạnh đó, hệ thống sử dụng thuật toán Bcrypt để băm mật khẩu, gắn token chống tấn công CSRF trên các form gửi dữ liệu và kiểm tra quyền sở hữu đơn hàng để ngăn chặn lỗi IDOR.")
add_para("Về chất lượng phần mềm, nhóm đặt mục tiêu xây dựng bộ kiểm thử tự động gồm 23 kịch bản kiểm thử bằng PHPUnit. Việc đặt hàng và trừ kho là khâu nhạy cảm đối với cửa hàng bán lẻ vì liên quan trực tiếp đến tiền bạc và hàng thực tế trong kho. Do đó, các kịch bản kiểm thử tự động giúp nhóm xác minh tính chính xác của các luồng xử lý trước khi hệ thống được đưa vào vận hành.")

add_sec("1.3. Phạm vi thực hiện")
add_para("Hệ thống Smart Casual được xây dựng để phục vụ nhu cầu mua sắm của khách hàng và công tác vận hành bán lẻ của nhân sự cửa hàng. Ứng dụng đảm nhiệm các nhóm công việc chính bao gồm trưng bày danh mục sản phẩm, quản lý giỏ hàng, thực hiện thanh toán trực tuyến, theo dõi tiến trình đơn hàng và thống kê kết quả kinh doanh.")
add_para("Đối với người mua sắm, hệ thống phân chia thành hai nhóm khách hàng với quyền hạn khác nhau. Khách hàng vãng lai có thể duyệt danh mục, tìm kiếm, lọc sản phẩm theo tiêu chí, chọn kích cỡ vào giỏ hàng lưu trong Session và đặt hàng mà không cần đăng ký. Với khách hàng thành viên, người dùng đăng ký tài khoản bằng email. Sau khi đăng nhập, giỏ hàng tạm được gộp vào bảng GIO_HANG trong cơ sở dữ liệu, cho phép người dùng xem lịch sử mua hàng tại route /account/orders, tự hủy đơn chờ duyệt, xác nhận đã nhận hàng và gửi đánh giá nhận xét cho từng mẫu trang phục.")
add_para("Đối với nội bộ cửa hàng, hệ thống thiết lập hai vai trò nhân sự đăng nhập qua cổng quản trị /admin/login. Nhân viên bán hàng có quyền xem danh sách đơn hàng, xem chi tiết địa chỉ người nhận, cập nhật các bước giao nhận, đánh dấu đơn đã thu tiền và duyệt hiển thị nhận xét của khách. Quản trị viên sở hữu toàn bộ quyền của nhân viên, đồng thời có phạm vi quản lý rộng hơn khi được độc quyền thêm, sửa, xóa danh mục, cập nhật giá bán và tồn kho từng size, ẩn sản phẩm và xem báo cáo doanh thu tài chính tại route /admin/revenue.")
add_para("Về phạm vi chức năng, hệ thống tổ chức danh mục hàng hóa trong bảng DANH_MUC và sản phẩm trong bảng SAN_PHAM. Mỗi mẫu thời trang liên kết với bảng CHI_TIET_SP để lưu các biến thể kích cỡ có đơn giá và tồn kho độc lập. Tệp ảnh sản phẩm được tải lên qua Laravel Storage và lưu thông tin tại bảng HINH_ANH_SP. Khi một mẫu hàng ngừng kinh doanh, quản trị viên sử dụng cờ TrangThai để ẩn sản phẩm khỏi giao diện bán hàng mà không xóa dòng dữ liệu, bảo vệ liên kết khóa ngoại với các đơn hàng cũ.")
add_para("Ở khâu tìm kiếm và chọn hàng, giao diện hỗ trợ tìm theo tên, lọc danh mục hoặc khoảng giá và sắp xếp giá. Danh sách sản phẩm được phân trang cố định 16 sản phẩm mỗi trang, sử dụng phương thức withQueryString() để giữ nguyên tham số lọc trên URL khi chuyển trang. Cơ chế giỏ hàng kép lưu trữ song song trên Session của khách vãng lai và bảng GIO_HANG của khách thành viên, tự động hợp nhất các món đồ khi người dùng đăng nhập tài khoản.")
add_para("Trong quy trình đặt hàng, khách hàng có thể chọn phương thức giao hàng và nhập mã giảm giá từ bảng MA_GIAM_GIA. Tại lớp dịch vụ CheckoutService, hệ thống sử dụng giao dịch kết hợp lệnh khóa dòng lockForUpdate trên bảng CHI_TIET_SP để trừ kho an toàn, ngăn chặn việc bán vượt số lượng thực tế. Khi đơn bị hủy theo ma trận ALLOWED_TRANSITIONS, hệ thống tự động hoàn kho và trả lại lượt dùng mã giảm giá. Dữ liệu từ các đơn hoàn tất được tổng hợp thành báo cáo doanh thu và giá trị đơn trung bình AOV theo từng tháng cho người quản lý.")
add_para("Bên cạnh các chức năng đã xây dựng, phiên bản hiện tại chưa bao gồm việc kết nối trực tiếp với cổng thanh toán điện tử, chưa có dịch vụ gửi email thông báo tự động sau khi đặt đơn, và chưa thiết lập tầng bộ nhớ đệm cho dữ liệu catalog sản phẩm. Các hạng mục này nằm ngoài phạm vi thực hiện của đồ án.")

# ----------------- CHƯƠNG 2 -----------------
add_chapter("CHƯƠNG 2. PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG")
add_para("Chương này trình bày chi tiết quá trình phân tích đối tượng người dùng và thiết kế kiến trúc dữ liệu cho website Smart Casual. Nội dung tập trung vào việc mô hình hóa các ca sử dụng qua sơ đồ Use Case, cấu trúc 13 bảng nghiệp vụ trong cơ sở dữ liệu quan hệ, các quyết định thiết kế biến thể sản phẩm và sự phối hợp giữa Database Trigger với Service Layer trong việc kiểm soát dữ liệu hóa đơn.")

add_sec("2.1. Phân tích đối tượng và vai trò người dùng")
add_para("Trong quá trình phân tích hệ thống thương mại điện tử bán lẻ thời trang, nhóm nhận thấy việc gom chung tài khoản quản trị và tài khoản khách hàng vào một bảng dữ liệu duy nhất thường dẫn đến nguy cơ nhầm lẫn quyền hạn hoặc phức tạp hóa cấu trúc bảng. Khách mua hàng cần các trường thông tin phục vụ giao nhận như địa chỉ nhà hay số điện thoại, trong khi nhân viên quản trị cần chức vụ và mã phòng ban. Việc phân tách rõ ràng các nhóm tác nhân giúp hệ thống kiểm soát quyền truy cập chặt chẽ ngay từ tầng định tuyến và tầng dữ liệu.")
add_para("Nhóm người dùng bên ngoài bao gồm khách hàng vãng lai và khách hàng thành viên. Khách hàng vãng lai chiếm phần lớn lưu lượng truy cập khi vào website tìm kiếm mẫu mã trang phục. Họ có quyền xem danh mục, lọc theo khoảng giá, xem chi tiết từng kích cỡ và thêm sản phẩm vào giỏ hàng phiên làm việc được lưu trong Session. Để giảm thiểu rào cản mua sắm, nhóm cho phép khách vãng lai tiến hành đặt hàng trực tiếp mà không bắt buộc tạo tài khoản. Khách hàng thành viên là người dùng đã đăng ký tài khoản bằng email. Khi đăng nhập, toàn bộ dữ liệu giỏ hàng tạm thời trong Session được hệ thống tự động gộp vào bảng GIO_HANG trong cơ sở dữ liệu. Khách hàng thành viên có trang quản lý cá nhân để theo dõi lịch sử các đơn hàng đã đặt, tự hủy các đơn đang ở trạng thái chờ duyệt, xác nhận đã nhận hàng khi kiện hàng đang giao và gửi đánh giá nhận xét cho từng mẫu trang phục.")
add_para("Nhóm người dùng nội bộ bao gồm nhân viên bán hàng và quản trị viên cửa hàng, đều đăng nhập tập trung qua cổng quản trị tại route /admin/login. Nhân viên bán hàng phụ trách xử lý đơn hàng phát sinh hằng ngày. Vai trò này có quyền tra cứu danh sách đơn, xem thông tin giao nhận, chuyển đổi trạng thái đơn qua các bước đóng gói hay xuất kho, đánh dấu đơn đã thu tiền và duyệt hiển thị nhận xét của khách hàng. Nhân viên không được cấp quyền chỉnh sửa giá bán sản phẩm, không được xóa danh mục và không được xem báo cáo tài chính của cửa hàng.")
add_para("Quản trị viên là người phụ trách cao nhất về mặt nghiệp vụ kinh doanh. Vai trò này kế thừa toàn bộ các quyền hạn xử lý đơn của nhân viên, đồng thời nắm giữ độc quyền các chức năng nhạy cảm gồm thêm mới, chỉnh sửa giá bán, cập nhật số lượng tồn kho từng size, ẩn sản phẩm khỏi trang bán hàng và truy cập báo cáo doanh thu tài chính. Việc thiết lập ranh giới phân quyền rõ rệt giữa hai vai trò nội bộ giúp ngăn chặn rủi ro nhân viên vô ý thay đổi giá bán hoặc xóa nhầm danh mục sản phẩm đang kinh doanh.")

add_sec("2.2. Sơ đồ UML Use Case tổng quan")
add_para("Mối quan hệ tương tác giữa 4 nhóm tác nhân và các chức năng chính của hệ thống Smart Casual được mô hình hóa như Hình 2.1 dưới đây.")
add_box("SƠ ĐỒ USE CASE HỆ THỐNG SMART CASUAL")
add_caption("Hình 2.1: Sơ đồ Use Case tổng quan thể hiện tương tác của từng vai trò trong hệ thống Smart Casual")
add_para("Từ sơ đồ Use Case, các ca sử dụng được nhóm phân bổ theo luồng nghiệp vụ thực tế của từng phân hệ: Phân hệ khách hàng tập trung vào luồng mua sắm trực tuyến từ xem catalog, chọn size, quản lý giỏ hàng đến đặt hàng. Khách thành viên mở rộng thêm quyền theo dõi lịch sử đơn, hủy đơn pending, xác nhận nhận hàng và đánh giá. Phân hệ quản trị tổ chức các ca sử dụng theo mô hình RBAC: nhân viên và quản trị viên đều xử lý đơn hàng, nhưng các ca sử dụng tác động trực tiếp đến danh mục, sản phẩm và báo cáo tài chính chỉ cho phép duy nhất quản trị viên thực hiện.")

add_sec("2.3. Sơ đồ và thiết kế Cơ sở dữ liệu (ERD)")
add_para("Cơ sở dữ liệu của hệ thống được xây dựng trên hệ quản trị MySQL 8.4, bao gồm 13 bảng nghiệp vụ liên kết chặt chẽ qua các ràng buộc khóa ngoại. Sơ đồ quan hệ thực thể của cơ sở dữ liệu được thể hiện tại Hình 2.2 dưới đây.")
add_box("SƠ ĐỒ QUAN HỆ THỰC THỂ (ERD) 13 BẢNG NGHIỆP VỤ")
add_caption("Hình 2.2: Sơ đồ quan hệ thực thể (ERD) của hệ thống thương mại điện tử Smart Casual")

add_para("Trong quá trình xây dựng sơ đồ cơ sở dữ liệu, nhóm đã đưa ra ba quyết định thiết kế then chốt:")
add_para("Thứ nhất là quyết định tách riêng bảng SAN_PHAM và bảng CHI_TIET_SP. Một mẫu áo sơ mi hay quần tây khi nhập về có các thuộc tính chung không đổi như tên gọi, mô tả chất liệu và nhóm danh mục. Tuy nhiên, mỗi sản phẩm lại có nhiều kích cỡ khác nhau, mỗi kích cỡ có số lượng tồn kho riêng biệt và có thể có đơn giá bán khác nhau giữa các size thông thường và size lớn. Nếu gom toàn bộ vào bảng SAN_PHAM, cơ sở dữ liệu không thể quản lý lượng tồn của từng size, dẫn đến việc không thể phát hiện tình trạng hết hàng khi khách chọn đúng size đó. Do đó, nhóm thiết kế bảng CHI_TIET_SP đóng vai trò là đơn vị lưu kho SKU, liên kết với SAN_PHAM qua khóa ngoại MaSP và bảng KICH_THUOC qua khóa ngoại MaKichThuoc. Ràng buộc UNIQUE(MaSP, MaKichThuoc) được thiết lập để đảm bảo mỗi sản phẩm chỉ liên kết với một kích cỡ duy nhất một lần.")
add_para("Thứ hai là việc lưu trữ đơn giá chụp tại thời điểm giao dịch trong bảng CHI_TIET_DON_HANG. Khi khách hàng đặt mua một món đồ, giá bán của biến thể tại thời điểm đó được sao chép vào cột DonGia của chi tiết đơn hàng. Giá bán niêm yết trong bảng CHI_TIET_SP có thể thay đổi trong tương lai do cửa hàng điều chỉnh giá thị trường hoặc áp dụng khuyến mãi mới. Nếu bảng CHI_TIET_DON_HANG chỉ lưu khóa ngoại MaCTSP mà không lưu lại giá bán thực tế lúc mua, tổng tiền của các hóa đơn cũ đã lưu trữ trong lịch sử sẽ bị sai lệch mỗi khi quản trị viên cập nhật giá sản phẩm.")
add_para("Thứ ba là việc sử dụng khóa chính phức hợp trên hai bảng trung gian GIO_HANG và CHI_TIET_DON_HANG. Bảng GIO_HANG sử dụng cặp (MaKH, MaCTSP) làm khóa chính, đảm bảo mỗi khách hàng chỉ có một dòng duy nhất cho một biến thể sản phẩm trong giỏ hàng. Khi khách hàng bấm thêm món đồ đã có sẵn trong giỏ, hệ thống tăng số lượng chứ không tạo thêm dòng mới gây trùng lặp. Tương tự, bảng CHI_TIET_DON_HANG sử dụng cặp (MaDonHang, MaCTSP) làm khóa chính, đảm bảo tính duy nhất của từng món hàng trong một hóa đơn.")

add_sec("2.4. Danh mục 13 bảng nghiệp vụ chi tiết")
add_para("Cấu trúc chi tiết của 13 bảng nghiệp vụ trong cơ sở dữ liệu:")
add_para("1. DANH_MUC: MaDanhMuc (VARCHAR(10), PK), TenDanhMuc (VARCHAR(100)), TrangThai (BOOLEAN). Bảng phân loại hàng hóa chính (Áo, Quần, Giày), dùng làm tiêu chí lọc và phân cấp danh mục.")
add_para("2. KICH_THUOC: MaKichThuoc (VARCHAR(10), PK), TenKichThuoc (VARCHAR(50)), TrangThai (BOOLEAN). Quy chuẩn hóa kích cỡ dùng chung (S, M, L, XL, 38, 39, 40...), tránh nhập tự do sai chính tả.")
add_para("3. SAN_PHAM: MaSP (VARCHAR(20), PK), MaDanhMuc (VARCHAR(10), FK), TenSanPham (VARCHAR(200)), MoTa (TEXT), TrangThai (BOOLEAN). Lưu thông tin chung. Trường TrangThai phục vụ xóa mềm logic để bảo vệ khóa ngoại với đơn hàng cũ.")
add_para("4. CHI_TIET_SP: MaCTSP (VARCHAR(30), PK), MaSP (VARCHAR(20), FK), MaKichThuoc (VARCHAR(10), FK), GiaBan (DECIMAL(18,2)), SoLuong (INT UNSIGNED), TrangThai (BOOLEAN). Ràng buộc UNIQUE(MaSP, MaKichThuoc). Đơn vị lưu kho SKU thực tế, là trọng tâm trong các truy vấn khóa dòng và trừ tồn kho khi checkout.")
add_para("5. HINH_ANH_SP: MaHinhAnh (VARCHAR(30), PK), MaSP (VARCHAR(20), FK), DuongDan (VARCHAR(255)), LaAnhChinh (BOOLEAN), TrangThai (BOOLEAN). Quản lý nhiều ảnh chụp cho mỗi sản phẩm.")
add_para("6. KHACH_HANG: MaKH (VARCHAR(20), PK), TenDangNhap (VARCHAR(100), UK), MatKhau (VARCHAR(255)), HoTen (VARCHAR(100)), SoDienThoai (VARCHAR(20)), Email (VARCHAR(100), UK), DiaChi (VARCHAR(255)), TrangThai (BOOLEAN). Phục vụ Guard web. Khóa chính là chuỗi ký tự MaKH thay vì số tự tăng id.")
add_para("7. NHAN_VIEN: MaNV (VARCHAR(20), PK), TenDangNhap (VARCHAR(100), UK), MatKhau (VARCHAR(255)), HoTen (VARCHAR(100)), ChucVu (VARCHAR(50): manager / staff), TrangThai (BOOLEAN). Phục vụ Guard admin. Phân tách rạch ròi với bảng KHACH_HANG để ngăn ngừa phân quyền nhầm lẫn.")
add_para("8. MA_GIAM_GIA: MaCode (VARCHAR(50), PK), GiaTri (DECIMAL(18,2)), LoaiGiamGia (VARCHAR(20): percent / cash / shipping), SoLuong (INT), TrangThai (BOOLEAN). Cột SoLuong bị trừ 1 khi đặt hàng thành công và được hoàn trả khi đơn hàng bị hủy trước khâu đóng gói.")
add_para("9. PHUONG_THUC_VAN_CHUYEN: MaPTVC (VARCHAR(20), PK), TenPTVC (VARCHAR(100)), GiaVanChuyen (DECIMAL(18,2)), TrangThai (BOOLEAN). Quy định cước phí cố định ('HT': 30.000 ₫, 'TC': 15.000 ₫) để áp tự động vào hóa đơn.")
add_para("10. GIO_HANG: MaKH (VARCHAR(20), FK), MaCTSP (VARCHAR(30), FK), SoLuong (INT). Khóa chính: PRIMARY KEY(MaKH, MaCTSP). Khóa chính phức hợp hai cột, đòi hỏi ghi đè setKeysForSaveQuery trong Eloquent Model.")
add_para("11. DON_HANG: MaDonHang (VARCHAR(30), PK), MaKH (FK, NULL), MaNV (FK, NULL), MaCode (FK, NULL), MaPTVC (FK), TenNguoiNhan (VARCHAR(100)), SoDienThoai (VARCHAR(20)), Email (VARCHAR(100)), NgayDat (DATETIME), DiaChiGiao (VARCHAR(255)), TongSoLuong (INT), TongTien (DECIMAL(18,2)), GiamGia (DECIMAL(18,2)), PhiVanChuyen (DECIMAL(18,2)), ThanhTien (DECIMAL(18,2)), PhuongThucThanhToan (VARCHAR(50)), DaThanhToan (BOOLEAN), TrangThaiDonHang (VARCHAR(50)). Lưu thông tin người nhận tại thời điểm đặt, tránh việc đổi địa chỉ tài khoản làm sai lệch đơn hàng cũ.")
add_para("12. CHI_TIET_DON_HANG: MaDonHang (VARCHAR(30), FK), MaCTSP (VARCHAR(30), FK), SoLuong (INT), DonGia (DECIMAL(18,2)), ThanhTien (DECIMAL(18,2)). Khóa chính: PRIMARY KEY(MaDonHang, MaCTSP). Quy tắc ON DELETE RESTRICT trên MaCTSP ngăn chặn xóa biến thể khi đã phát sinh đơn hàng lịch sử.")
add_para("13. DANH_GIA: MaDanhGia (VARCHAR(30), PK), MaSP (VARCHAR(20), FK), MaKH (VARCHAR(20), FK, NULL), TenKhachHang (VARCHAR(100)), SoSao (TINYINT), NoiDung (TEXT), TrangThai (BOOLEAN). Trường TrangThai giúp nhân viên kiểm duyệt ẩn các nhận xét vi phạm hoặc spam trước khi công khai.")

add_sec("2.5. Thiết kế Triggers và tính toán tự động")
add_para("Trong các hệ thống thương mại điện tử, việc tính toán tổng số tiền hóa đơn nếu chỉ giao phó cho code ứng dụng ở Controller sẽ dễ dẫn đến sai lệch khi có nhiều luồng can thiệp vào cơ sở dữ liệu. Để giải quyết bài toán này, nhóm phân chia trách nhiệm tính toán thành hai tầng kiểm soát như mô tả tại Hình 2.3 dưới đây.")
add_box("SƠ ĐỒ PHỐI HỢP TRÁCH NHIỆM GIỮA SERVICE LAYER VÀ DATABASE TRIGGERS")
add_caption("Hình 2.3: Sơ đồ phối hợp trách nhiệm giữa Service Layer của ứng dụng và Database Triggers trong MySQL")

add_para("Tầng Database Triggers trong MySQL chịu trách nhiệm bảo vệ tính bất biến toán học ở mức lưu trữ dữ liệu, gồm 4 triggers được khai báo qua Migration: Trigger trg_tudong_lay_phiship_insert kích hoạt trước khi chèn một bản ghi vào DON_HANG. Nếu trường cước phí vận chuyển chưa có hoặc bằng 0, trigger tự động tra cứu mức cước quy định trong bảng PHUONG_THUC_VAN_CHUYEN theo mã MaPTVC đã chọn để gán vào đơn hàng. Đồng thời, trigger áp dụng công thức tính toán số tiền thực thu là NEW.ThanhTien = GREATEST(0, NEW.TongTien + NEW.PhiVanChuyen - NEW.GiamGia). Hàm GREATEST(0, ...) đảm bảo tổng số tiền thanh toán không bao giờ nhận giá trị âm khi mức giảm giá vượt quá tiền hàng và cước ship.")
add_para("Ba trigger còn lại xử lý sự kiện trên bảng CHI_TIET_DON_HANG. Trigger trg_capnhat_tong_donhang_after_insert tự động thực hiện hàm SUM trên các dòng chi tiết để cập nhật lại TongSoLuong và TongTien vào bảng DON_HANG ngay khi có món hàng mới được thêm. Trigger trg_capnhat_tong_donhang_after_delete thực hiện phép trừ tương ứng khi một dòng chi tiết bị xóa khỏi đơn. Riêng trigger trg_capnhat_tong_donhang_after_update xử lý khi số lượng hoặc đơn giá của một dòng chi tiết bị sửa đổi, tự động cập nhật lại tổng tiền cho cả đơn hàng cũ và đơn hàng mới nếu có sự thay đổi mã đơn.")
add_para("Tầng Application Service (CheckoutService) chịu trách nhiệm kiểm soát các quy tắc nghiệp vụ phức tạp mà trigger không thể can thiệp. Service thực hiện kiểm tra tính hợp lệ của mã giảm giá thông qua phương thức isValid(), tính toán số tiền khấu trừ tương ứng với từng loại voucher, và bọc toàn bộ chuỗi thao tác trong một giao dịch DB::transaction(). Khi đọc tồn kho của từng biến thể sản phẩm, Service sử dụng lệnh lockForUpdate() để khóa dòng dữ liệu, buộc các request đặt hàng đồng thời cho cùng một sản phẩm phải xếp hàng đợi, ngăn ngừa hiện tượng bán vượt số lượng tồn kho. Khi đơn hàng bị hủy, hàm cancelOrder() trong Service kiểm tra ma trận trạng thái ALLOWED_TRANSITIONS, sau đó tự động hoàn trả số lượng vào bảng CHI_TIET_SP và cộng lại lượt dùng cho MA_GIAM_GIA.")
add_para("Việc kết hợp hai tầng giúp dữ liệu trong cơ sở dữ liệu luôn chính xác ngay cả khi được cập nhật trực tiếp qua các câu lệnh SQL ngoài ứng dụng. Tuy nhiên, cách làm này dẫn đến sự đánh đổi kỹ thuật khi logic tính tiền nằm ở cả hai nơi gồm file SQL tạo trigger và class CheckoutService trong mã nguồn PHP, đòi hỏi người bảo trì phải đồng bộ cả hai tầng mỗi khi có thay đổi về quy tắc tính phí giao nhận hay làm tròn số tiền.")

# ----------------- CHƯƠNG 3 -----------------
add_chapter("CHƯƠNG 3. CÔNG NGHỆ, KIẾN TRÚC VÀ TRIỂN KHAI")
add_para("Chương này trình bày chi tiết về kiến trúc phần mềm, danh mục công nghệ và giải pháp kỹ thuật triển khai cho website Smart Casual. Nội dung tập trung phân tích luồng xử lý request qua các tầng kiến trúc, cơ chế xác thực đa guard, kỹ thuật xử lý khóa chính phức hợp trong Eloquent ORM, giải pháp khóa dòng dữ liệu chống tranh chấp tồn kho và hướng dẫn cài đặt hệ thống từ mã nguồn.")

add_sec("3.1. Kiến trúc tổng thể hệ thống")
add_para("Hệ thống được tổ chức theo kiến trúc Model - View - Controller (MVC) mở rộng với lớp dịch vụ (Service Layer) của framework Laravel 11.x. Sơ đồ luồng xử lý HTTP request qua các tầng kiến trúc của hệ thống được thể hiện như Hình 3.1 dưới đây.")
add_box("SƠ ĐỒ KIẾN TRÚC VÀ LUỒNG XỬ LÝ REQUEST CỦA HỆ THỐNG")
add_caption("Hình 3.1: Sơ đồ luồng xử lý HTTP request qua các tầng kiến trúc của hệ thống Smart Casual")

add_para("Một request từ phía trình duyệt người dùng được xử lý tuần tự qua các tầng thành phần: public/index.php tiếp nhận request -> Router (routes/web.php) định tuyến -> Middleware Pipeline kiểm tra phiên và phân quyền (AdminMiddleware, ManagerOnlyMiddleware, throttle:login) -> Form Request xác thực dữ liệu đầu vào -> Controller ủy quyền cho Service Layer (CartService, CheckoutService) xử lý nghiệp vụ phức tạp -> Eloquent Models truy vấn xuống MySQL 8.4 qua PDO Parameter Binding an toàn -> Controller truyền kết quả sang View Blade để render giao diện HTML gửi về cho người dùng.")

add_sec("3.2. Danh mục công nghệ và thư viện sử dụng")
add_para("Danh mục công nghệ được nhóm lựa chọn gồm: PHP phiên bản 8.3 [3] tận dụng typed properties, match expressions, constructor property promotion; Laravel 11.x, phiên bản ổn định tại thời điểm nhóm khởi tạo dự án [2], cung cấp cấu trúc MVC, Eloquent ORM, Routing và Migrations; MySQL Community Server 8.4.3 [4] hỗ trợ InnoDB, khóa ngoại, ACID transactions, lockForUpdate và Database Triggers; Blade Templating Engine kết hợp Bootstrap 5.2.3 [5] và FontAwesome 6 Free; PHPUnit 10.5.66 [6] thực thi bộ 23 kịch bản kiểm thử tự động trên cơ sở dữ liệu độc lập; Composer 2.9.4 quản lý các gói phụ thuộc theo chuẩn PSR-4.")

add_sec("3.3. Cơ chế xác thực và phân quyền (Multi-Auth & Middleware)")
add_para("Hệ thống cấu hình hai Guard độc lập trong config/auth.php: Guard web sử dụng provider khach_hangs trỏ vào Model KhachHang (bảng KHACH_HANG, khóa chính MaKH, mật khẩu MatKhau); Guard admin sử dụng provider nhan_viens trỏ vào Model NhanVien (bảng NHAN_VIEN, khóa chính MaNV, mật khẩu MatKhau). Cả hai Model đều cài đặt giao diện Authenticatable và ghi đè getAuthPasswordName() để trỏ vào cột MatKhau. Việc phân tách giúp phiên đăng nhập của nhân viên và khách hàng hoàn toàn tách biệt, không bị ghi đè Session trên cùng một trình duyệt.")
add_para("Hệ thống phân quyền được thực thi qua Middleware: AdminMiddleware bảo vệ toàn bộ route /admin (trừ /admin/login); ManagerOnlyMiddleware bảo vệ route doanh thu /admin/revenue và các thao tác thêm, sửa, xóa sản phẩm hoặc danh mục. Route xem chi tiết đơn hàng /account/orders/{id} được ràng buộc điều kiện MaKH = Auth::guard('web')->user()->MaKH, trả về mã lỗi HTTP 403 Forbidden nếu khách hàng cố tình truy cập đơn của người khác, ngăn chặn hoàn toàn lỗ hổng IDOR. Các route đăng nhập đều được gắn throttle:login giới hạn tối đa 5 lần thử trong một phút để chống dò quét mật khẩu.")

add_sec("3.4. Mô hình ORM Eloquent và bảo vệ toàn vẹn dữ liệu")
add_para("Đối với khóa chính phức hợp trên bảng GIO_HANG (MaKH, MaCTSP), nhóm đã ghi đè phương thức setKeysForSaveQuery trong Model GioHang để ép buộc Eloquent đưa đủ cả hai điều kiện WHERE MaKH = ? AND MaCTSP = ? vào câu lệnh cập nhật, ngăn ngừa việc ghi đè lên các sản phẩm khác của cùng một người dùng.")
add_para("Đối với tranh chấp tồn kho khi đặt hàng, nhóm sử dụng lệnh lockForUpdate() bên trong DB::transaction() tại CheckoutService. Khi hai khách hàng cùng bấm thanh toán cho sản phẩm chỉ còn tồn kho bằng 1, lệnh này khóa dòng dữ liệu tại mức CSDL InnoDB, buộc request thứ hai phải chờ cho đến khi giao dịch thứ nhất hoàn tất. Lúc request thứ hai đọc lại dữ liệu thì tồn kho đã bằng 0 và hệ thống sẽ báo lỗi hết hàng, ngăn chặn hoàn toàn hiện tượng bán vượt tồn kho Overselling. Hệ thống cũng áp dụng mã idempotency_key gửi kèm từ client để chống double-submit khi khách bấm đặt hàng nhiều lần. Toàn bộ câu lệnh đều sử dụng Parameter Binding chống SQL Injection, token @csrf chống CSRF và cú pháp {{ }} của Blade chống XSS.")

add_sec("3.5. Yêu cầu môi trường hệ thống")
add_para("Hệ thống yêu cầu: HĐH Windows 10/11, Linux hoặc macOS; PHP >= 8.2 (khuyên dùng PHP 8.3 với đầy đủ extension pdo_mysql, mbstring, openssl, tokenizer, xml, ctype, json, curl, fileinfo); Composer >= 2.2; MySQL 8.x chạy trên cổng 3306 mặc định.")

add_sec("3.6. Các bước cài đặt chi tiết từ mã nguồn")
add_para("Quy trình cài đặt gồm 6 bước: Bước 1 chuyển vào thư mục cd SMART_CASUAL_LARAVEL; Bước 2 chạy composer install --prefer-dist --no-interaction; Bước 3 tạo file cấu hình copy .env.example .env và khai báo thông tin DB_DATABASE=smart_casual_shop; Bước 4 sinh mã khóa php artisan key:generate và tạo symlink php artisan storage:link; Bước 5 tạo database trong MySQL (CREATE DATABASE smart_casual_shop CHARACTER SET utf8mb4;) rồi chạy php artisan migrate:fresh --seed (hoặc import file database/smart_casual_shop_dump.sql); Bước 6 khởi chạy máy chủ php artisan serve và truy cập http://localhost:8000.")
add_para("Một số lỗi thường gặp khi cài đặt: lỗi thiếu khóa ứng dụng (xử lý bằng php artisan key:generate); lỗi chưa tạo database trước khi migrate (xử lý bằng cách tạo database trước); lỗi ảnh sản phẩm 404 (xử lý bằng php artisan storage:link); lỗi thiếu file fixture seeder (xử lý bằng cách import trực tiếp file dump SQL dự phòng).")

add_sec("3.7. Danh sách tài khoản thử nghiệm hệ thống")
add_para("Danh sách tài khoản thử nghiệm được trình bày tại Bảng 3.1 dưới đây.")
add_tbl_caption("Bảng 3.1: Danh sách tài khoản thử nghiệm hệ thống Smart Casual")

demo_tbl = doc.add_table(rows=5, cols=4)
demo_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(demo_tbl)

d_headers = ["Vai trò (Role)", "Cổng đăng nhập", "Tài khoản / Email", "Mật khẩu"]
for i, h in enumerate(d_headers):
    c = demo_tbl.rows[0].cells[i]
    set_cell_background(c, "14221D")
    set_cell_margins(c, 60, 60, 80, 80)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

demo_rows = [
    ("Khách hàng 1 (Customer)", "/login", "demo1@gmail.com", "password123"),
    ("Khách hàng 2 (IDOR Test)", "/login", "demo2@gmail.com", "password123"),
    ("Nhân viên (Staff)", "/admin/login", "staff", "staff123"),
    ("Quản trị (Manager)", "/admin/login", "admin", "admin123")
]

for r_idx, (r0, r1, r2, r3) in enumerate(demo_rows, start=1):
    row = demo_tbl.rows[r_idx]
    bg = "F9FBFA" if r_idx % 2 == 1 else "FFFFFF"
    for c_idx, val in enumerate([r0, r1, r2, r3]):
        c = row.cells[c_idx]
        set_cell_background(c, bg)
        set_cell_margins(c, 50, 50, 60, 60)
        p = c.paragraphs[0]
        if c_idx in [1, 3]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(val)
        r.font.size = Pt(10)

# ----------------- CHƯƠNG 4 -----------------
add_chapter("CHƯƠNG 4. KẾT QUẢ VÀ KIỂM THỬ HỆ THỐNG")
add_para("Chương này trình bày toàn bộ kết quả giao diện và chức năng của hệ thống Smart Casual trên hai phân hệ Khách hàng và Quản trị viên, đối chiếu mức độ đáp ứng kỹ thuật so với đề cương học phần, đồng thời phân tích phương pháp và kết quả thực thi của bộ kiểm thử tự động gồm 23 kịch bản kiểm thử nhằm chứng minh độ tin cậy của mã nguồn.")

add_sec("4.1. Phân hệ Khách hàng (Customer Storefront)")
add_para("Quy trình mua sắm của khách hàng được xây dựng bám sát hành vi thực tế: Khách hàng vào trang chủ (Route /) xem banner và danh mục 48 sản phẩm thời trang cùng 4 mẫu bán chạy. Khách sử dụng bộ lọc tìm kiếm theo tên, danh mục (Áo, Quần, Giày), khoảng giá và sắp xếp tăng/giảm. Hệ thống phân trang 16 sản phẩm/trang và giữ lại tham số lọc trên URL qua withQueryString(). Giao diện trang chủ và bộ lọc catalog sản phẩm được thể hiện tại Hình 4.1.")
add_box("HÌNH 4.1: Giao diện trang chủ và bộ lọc catalog sản phẩm tại route /")
add_caption("Hình 4.1: Giao diện trang chủ và bộ lọc catalog sản phẩm tại route /")

add_para("Khách bấm vào sản phẩm để sang trang Chi tiết (Route /products/{id}), xem ảnh thumbnail chuyển đổi bằng JavaScript, bấm chọn kích cỡ (Size) để cập nhật giá và tồn kho khả dụng tương ứng. Nút chọn size hết hàng sẽ bị vô hiệu hóa. Khách xem điểm đánh giá sao, đọc nhận xét cũ và có thể gửi đánh giá mới. Giao diện chi tiết sản phẩm và lựa chọn biến thể kích cỡ được thể hiện tại Hình 4.2.")
add_box("HÌNH 4.2: Giao diện chi tiết sản phẩm và lựa chọn biến thể kích cỡ tại route /products/{id}")
add_caption("Hình 4.2: Giao diện chi tiết sản phẩm và lựa chọn biến thể kích cỡ tại route /products/{id}")

add_para("Tại trang Giỏ hàng (Route /cart), khách điều chỉnh số lượng (hệ thống tự động kiểm tra trần tồn kho qua CartService), xóa món hoặc làm trống giỏ hàng. Khi khách vãng lai đăng nhập, hệ thống tự động gộp giỏ hàng Session vào CSDL qua mergeGuestCartToCustomer. Giao diện giỏ hàng và cập nhật số lượng được thể hiện tại Hình 4.3.")
add_box("HÌNH 4.3: Giao diện giỏ hàng và cập nhật số lượng tại route /cart")
add_caption("Hình 4.3: Giao diện giỏ hàng và cập nhật số lượng tại route /cart")

add_para("Tại trang Thanh toán (Route /checkout), khách nhập thông tin giao hàng (tự điền nếu là thành viên), chọn cước vận chuyển Hỏa tốc (30.000 ₫) hoặc Tiêu chuẩn (15.000 ₫), nhập voucher giảm giá (SAVE10, FREESHIP), chọn hình thức COD hoặc Chuyển khoản và bấm Đặt hàng. Hệ thống áp dụng idempotency token và vô hiệu hóa nút submit để chống đặt trùng đơn. Giao diện thanh toán đơn hàng được thể hiện tại Hình 4.4.")
add_box("HÌNH 4.4: Giao diện thanh toán đơn hàng tại route /checkout")
add_caption("Hình 4.4: Giao diện thanh toán đơn hàng tại route /checkout")

add_para("Tại trang Lịch sử đơn hàng (Route /account/orders và /account/orders/{id}), khách xem tiến trình đơn hàng cá nhân, bấm 'Hủy đơn' nếu đơn còn ở trạng thái pending (hệ thống tự động hoàn kho), hoặc bấm 'Xác nhận đã nhận hàng' khi đơn đang shipping. Cơ chế kiểm tra quyền sở hữu ngăn chặn hoàn toàn lỗi IDOR. Giao diện lịch sử và chi tiết đơn hàng của khách hàng được thể hiện tại Hình 4.5.")
add_box("HÌNH 4.5: Giao diện lịch sử và chi tiết đơn hàng của khách hàng tại route /account/orders")
add_caption("Hình 4.5: Giao diện lịch sử và chi tiết đơn hàng của khách hàng tại route /account/orders")

add_sec("4.2. Phân hệ Quản trị viên (Admin Portal)")
add_para("Khu vực quản trị tại tiền tố /admin phân quyền rành mạch giữa Staff và Manager: Bảng điều khiển (Route /admin) cung cấp 5 thẻ thống kê doanh thu thực tế, tổng đơn, đơn chờ, tổng sản phẩm và số lượng thành viên, đi kèm danh sách 5 đơn hàng mới nhất và đánh giá vừa gửi. Bảng điều khiển tổng quan Admin Dashboard được thể hiện tại Hình 4.6.")
add_box("HÌNH 4.6: Bảng điều khiển tổng quan Admin Dashboard tại route /admin")
add_caption("Hình 4.6: Bảng điều khiển tổng quan Admin Dashboard tại route /admin")

add_para("Quản lý Sản phẩm (Route /admin/products) cho phép tìm kiếm, thêm sản phẩm mới có upload ảnh qua Storage (/admin/products/create), chỉnh sửa giá/kho từng size (/admin/products/{id}/edit), ẩn sản phẩm (xóa mềm logic) và khôi phục hoạt động kinh doanh mà không ảnh hưởng đơn hàng cũ. Giao diện quản lý sản phẩm và form thêm mới được thể hiện tại Hình 4.7.")
add_box("HÌNH 4.7: Giao diện quản lý sản phẩm và form thêm mới tại route /admin/products")
add_caption("Hình 4.7: Giao diện quản lý sản phẩm và form thêm mới tại route /admin/products")

add_para("Quản lý Danh mục (Route /admin/categories) hiển thị số lượng sản phẩm liên kết của từng danh mục và chặn thao tác xóa nếu danh mục đang chứa sản phẩm con. Giao diện quản lý danh mục sản phẩm được thể hiện tại Hình 4.8.")
add_box("HÌNH 4.8: Giao diện quản lý danh mục sản phẩm tại route /admin/categories")
add_caption("Hình 4.8: Giao diện quản lý danh mục sản phẩm tại route /admin/categories")

add_para("Quản lý Đơn hàng (Route /admin/orders và /admin/orders/{id}) cho phép lọc theo trạng thái quy trình, cập nhật tiến trình giao nhận theo ma trận ALLOWED_TRANSITIONS, xác nhận thanh toán tiền hàng và tự động hoàn trả tồn kho nếu hủy đơn. Giao diện chi tiết đơn hàng và cập nhật trạng thái được thể hiện tại Hình 4.9.")
add_box("HÌNH 4.9: Giao diện chi tiết đơn hàng và cập nhật trạng thái tại route /admin/orders/{id}")
add_caption("Hình 4.9: Giao diện chi tiết đơn hàng và cập nhật trạng thái tại route /admin/orders/{id}")

add_para("Báo cáo Doanh thu (Route /admin/revenue) dành riêng cho Manager, hiển thị tổng doanh thu từ đơn delivered, giá trị đơn trung bình AOV, bảng thống kê doanh thu theo tháng và cơ cấu phương thức thanh toán. Giao diện báo cáo thống kê doanh thu được thể hiện tại Hình 4.10.")
add_box("HÌNH 4.10: Giao diện báo cáo thống kê doanh thu tại route /admin/revenue")
add_caption("Hình 4.10: Giao diện báo cáo thống kê doanh thu tại route /admin/revenue")

add_para("Kiểm duyệt Đánh giá (Route /admin/reviews) cho phép duyệt bật/tắt hiển thị công khai nhận xét của khách hàng hoặc xóa bỏ các bình luận spam. Giao diện kiểm duyệt đánh giá khách hàng được thể hiện tại Hình 4.11.")
add_box("HÌNH 4.11: Giao diện kiểm duyệt đánh giá khách hàng tại route /admin/reviews")
add_caption("Hình 4.11: Giao diện kiểm duyệt đánh giá khách hàng tại route /admin/reviews")

add_sec("4.3. Bảng đối chiếu yêu cầu kỹ thuật (Traceability Matrix)")
add_para("Bảng đối chiếu yêu cầu kỹ thuật được thể hiện tại Bảng 4.1 dưới đây.")
add_tbl_caption("Bảng 4.1: Bảng đối chiếu yêu cầu kỹ thuật trong đề cương môn học (Traceability Matrix)")

trace_tbl = doc.add_table(rows=17, cols=4)
trace_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(trace_tbl)

t_hdrs = ["STT", "Yêu cầu trong đề cương", "Triển khai thực tế trong mã nguồn", "Đánh giá"]
for i, h in enumerate(t_hdrs):
    c = trace_tbl.rows[0].cells[i]
    set_cell_background(c, "14221D")
    set_cell_margins(c, 50, 50, 60, 60)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

trace_data = [
    ("1", "Framework Laravel", "Laravel 11.x trên nền tảng ngôn ngữ PHP 8.3.", "ĐẠT"),
    ("2", "Cơ sở dữ liệu MySQL", "MySQL Community Server 8.4.3 với 13 bảng nghiệp vụ chuẩn hóa 3NF.", "ĐẠT"),
    ("3", "Khai báo cấu trúc bảng Migration", "4 tệp Migration khởi tạo 13 bảng, ràng buộc khóa ngoại và 4 Triggers.", "ĐẠT"),
    ("4", "Thao tác qua Eloquent ORM", "13 Models Eloquent đầy đủ quan hệ, Accessors và hỗ trợ Composite Key.", "ĐẠT"),
    ("5", "Seeder tạo dữ liệu mẫu", "ShopQuanAoSeeder nạp 841 bản ghi sạch; có file dự phòng dump SQL.", "ĐẠT"),
    ("6", "Giao diện Blade + Bootstrap", "Hệ thống Blade Templates kết hợp Bootstrap 5.2.3 và FontAwesome 6.", "ĐẠT"),
    ("7", "Tách riêng Admin và User", "Hai layout độc lập: layouts/app.blade.php và layouts/admin.blade.php.", "ĐẠT"),
    ("8", "Tối thiểu 2 vai trò người dùng", "3 vai trò Multi-Auth: Customer, Staff và Manager.", "ĐẠT"),
    ("9", "Đăng nhập / Đăng xuất Laravel Auth", "Multi-Auth Guards (web và admin), băm mật khẩu Bcrypt.", "ĐẠT"),
    ("10", "Phân quyền qua Middleware", "AdminMiddleware và ManagerOnlyMiddleware bảo vệ toàn diện /admin.", "ĐẠT"),
    ("11", "CRUD trên bảng có quan hệ", "CRUD hoàn chỉnh trên DANH_MUC và SAN_PHAM / CHI_TIET_SP, có chặn xóa.", "ĐẠT"),
    ("12", "Tìm kiếm kết hợp phân trang", "Tìm kiếm tên, lọc danh mục/giá, phân trang giữ query string withQueryString().", "ĐẠT"),
    ("13", "Validation dữ liệu nhập form", "Form Request: StoreProductRequest, CheckoutRequest... có báo lỗi tiếng Việt.", "ĐẠT"),
    ("14", "Thông báo Flash Message", "Flash session dismissible alert (success, error, info) trên cả hai layout.", "ĐẠT"),
    ("15", "Mở rộng: Upload file Storage", "Upload ảnh sản phẩm qua Storage::disk('public') kết hợp symlink.", "ĐẠT"),
    ("16", "Mở rộng: Xóa mềm / Ẩn sản phẩm", "Ẩn/khôi phục sản phẩm qua TrangThai để bảo toàn khóa ngoại đơn hàng cũ.", "ĐẠT")
]

for r_idx, (c0, c1, c2, c3) in enumerate(trace_data, start=1):
    row = trace_tbl.rows[r_idx]
    bg = "F9FBFA" if r_idx % 2 == 1 else "FFFFFF"
    for c_idx, val in enumerate([c0, c1, c2, c3]):
        c = row.cells[c_idx]
        set_cell_background(c, bg)
        set_cell_margins(c, 40, 40, 50, 50)
        p = c.paragraphs[0]
        if c_idx in [0, 3]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(val)
        if c_idx == 3:
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x00, 0x64, 0x00)
        r.font.size = Pt(9.5)

add_sec("4.4. Phương pháp kiểm thử tự động (Feature Tests)")
add_para("Nhóm áp dụng phương pháp kiểm thử tích hợp tự động bằng PHPUnit (php artisan test) trên cơ sở dữ liệu thử nghiệm riêng biệt smart_casual_test. Trait DatabaseTransactions được gắn vào toàn bộ 4 file Feature Tests để tự động rollback dữ liệu sau mỗi hàm test, không để lại dữ liệu rác trong database.")
add_para("Tổng cộng hệ thống có 23 kịch bản kiểm thử tự động gồm 1 test Unit mẫu kiểm tra môi trường và 22 test Feature nghiệp vụ bao phủ toàn bộ các luồng quan trọng.")

add_sec("4.5. Kết quả kiểm thử tự động")
add_para("Chi tiết các kịch bản kiểm thử:")
add_para("• AdminCrudAndMiddlewareTest (5 test): Kiểm tra khách vãng lai và khách thường bị chặn vào /admin; kiểm tra nhân viên staff đăng nhập thành công; kiểm tra CRUD danh mục và chặn xóa khi có sản phẩm con; kiểm tra thêm sản phẩm và xóa mềm logic.")
add_para("• AuthenticationAndAuthorizationTest (3 test): Kiểm tra đăng ký tài khoản có băm mật khẩu Bcrypt; kiểm tra đăng nhập và đăng xuất; kiểm tra bảo mật IDOR ngăn chặn khách hàng A xem trộm đơn hàng của khách hàng B (HTTP 403 Forbidden).")
add_para("• CartAndCheckoutWorkflowTest (7 test): Kiểm tra thêm/sửa/xóa giỏ hàng; kiểm tra đặt hàng nguyên tử và trừ tồn kho; kiểm tra chặn đặt hàng khi vượt quá tồn kho; kiểm tra composite key trên bảng GIO_HANG không bị ghi đè chéo giữa các mặt hàng hoặc giữa các thành viên; kiểm tra gộp giỏ hàng vãng lai khi đăng nhập; kiểm tra Database Trigger tính lại tổng tiền cho cả đơn cũ và mới.")
add_para("• CatalogAndProductTest (7 test): Kiểm tra quan hệ Eloquent; kiểm tra tải trang chủ; kiểm tra tìm kiếm theo từ khóa; kiểm tra lọc theo danh mục; kiểm tra phân trang giữ query string; kiểm tra chi tiết sản phẩm và gửi đánh giá nhận xét.")

add_para("Log thực thi thực tế khi chạy lệnh php artisan test trong Terminal:")
p_code = doc.add_paragraph()
p_code.paragraph_format.left_indent = Inches(0.2)
p_code.paragraph_format.space_before = Pt(4)
p_code.paragraph_format.space_after = Pt(6)
r_c = p_code.add_run("""   PASS  Tests\\Unit\\ExampleTest
  ✓ that true is true                                                                         0.01s

   PASS  Tests\\Feature\\AdminCrudAndMiddlewareTest
  ✓ guest cannot access admin dashboard                                                      0.20s
  ✓ customer cannot access admin dashboard                                                   0.03s
  ✓ staff can login and access admin dashboard                                               0.04s
  ✓ admin category crud workflow                                                             0.06s
  ✓ admin product crud and soft delete                                                       0.06s

   PASS  Tests\\Feature\\AuthenticationAndAuthorizationTest
  ✓ customer registration with password hashing                                              0.03s
  ✓ customer login and logout                                                                0.23s
  ✓ user cannot access other users order idor protection                                     0.05s

   PASS  Tests\\Feature\\CartAndCheckoutWorkflowTest
  ✓ cart add update remove operations                                                        0.04s
  ✓ checkout atomic workflow and stock deduction                                             0.03s
  ✓ cannot checkout more than available stock                                                0.03s
  ✓ member composite cart with three items update and delete                                 0.04s
  ✓ two members with same variant do not interfere                                           0.04s
  ✓ guest cart and member cart merge cleanly                                                 0.05s
  ✓ trigger update order items recalculates both old and new orders                          0.05s

   PASS  Tests\\Feature\\CatalogAndProductTest
  ✓ database relationships between category and product                                      0.03s
  ✓ storefront home page loads with products                                                 0.05s
  ✓ search products by keyword                                                               0.04s
  ✓ filter products by category                                                              0.06s
  ✓ pagination and query string retention                                                    0.06s
  ✓ product detail page loads                                                                0.04s
  ✓ submit product review                                                                    0.04s

  Tests:    23 passed (118 assertions)
  Duration: 1.42s""")
r_c.font.name = 'Consolas'
r_c.font.size = Pt(9.5)
r_c.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

add_para("Phân tích kết quả: Toàn bộ 23/23 kịch bản kiểm thử đều đạt kết quả PASS với 118 assertions trong 1.42 giây. Kết quả này chứng minh hệ thống phân quyền chặn đúng đối tượng, các giao dịch tài chính và trừ kho được thực hiện nguyên tử an toàn, và các triggers tự động hóa cập nhật số liệu chính xác.")

# ----------------- CHƯƠNG 5 -----------------
add_chapter("CHƯƠNG 5. ĐÁNH GIÁ NHÓM, HẠN CHẾ VÀ HƯỚNG PHÁT TRIỂN")
add_para("Chương này tổng kết quá trình tổ chức thực hiện đề tài thông qua bảng phân công nhiệm vụ chi tiết và phương thức phối hợp kỹ thuật của các thành viên trong nhóm. Đồng thời, chương thẳng thắn chỉ ra những hạn chế kỹ thuật còn tồn tại trong phiên bản hiện tại và đề ra lộ trình nâng cấp hệ thống theo 4 giai đoạn cụ thể.")

add_sec("5.1. Bảng phân công nhiệm vụ chi tiết")
add_para("Bảng phân công nhiệm vụ chi tiết của các thành viên trong nhóm được thể hiện tại Bảng 5.1 dưới đây.")
add_tbl_caption("Bảng 5.1: Bảng phân công nhiệm vụ chi tiết của các thành viên trong nhóm")

assign_tbl = doc.add_table(rows=4, cols=5)
assign_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(assign_tbl)

a_hdrs = ["STT", "Họ và tên", "MSSV - Lớp", "Nhiệm vụ chính đảm nhiệm", "Mức độ hoàn thành"]
for i, h in enumerate(a_hdrs):
    c = assign_tbl.rows[0].cells[i]
    set_cell_background(c, "14221D")
    set_cell_margins(c, 50, 50, 60, 60)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

assign_data = [
    ("1", "Nguyễn Huy Hoàng", "2001230266\n14DHTH02", "Khởi tạo dự án Laravel 11.x; viết 4 tệp Migration khởi tạo 13 bảng và 4 Database Triggers; xây dựng 13 Eloquent Models kèm quan hệ; triển khai lớp dịch vụ CheckoutService (lockForUpdate) và CartService; lập trình các Controller khu vực quản trị; thiết kế và viết trọn bộ 23 kịch bản kiểm thử tự động với PHPUnit.", "100%"),
    ("2", "Nguyễn Việt Vương", "2001231062\n14DHTH02", "Chuyển đổi giao diện tĩnh HTML sang hệ thống Blade Templates; xây dựng hai master layout layouts/app.blade.php và layouts/admin.blade.php; thiết kế giao diện các trang Storefront gồm Catalog, Chi tiết sản phẩm, Giỏ hàng và Thanh toán; kiểm tra độ tương thích trên thiết bị di động qua Bootstrap 5.2.3; lập trình các đoạn mã JavaScript phía client.", "100%"),
    ("3", "Nguyễn Ngọc Tường Vi", "2001231032\n14DHTH12", "Thu thập, chuẩn hóa dữ liệu mẫu và lập trình seeder ShopQuanAoSeeder.php nạp 841 bản ghi vào CSDL; lập trình chức năng đăng ký, đăng nhập tài khoản khách hàng; xây dựng cấu trúc báo cáo đồ án; chuẩn bị tài liệu hướng dẫn cài đặt README.md, xuất file kết xuất CSDL dump SQL và slide thuyết trình.", "100%")
]

for r_idx, (c0, c1, c2, c3, c4) in enumerate(assign_data, start=1):
    row = assign_tbl.rows[r_idx]
    bg = "F9FBFA" if r_idx % 2 == 1 else "FFFFFF"
    for c_idx, val in enumerate([c0, c1, c2, c3, c4]):
        c = row.cells[c_idx]
        set_cell_background(c, bg)
        set_cell_margins(c, 40, 40, 50, 50)
        p = c.paragraphs[0]
        if c_idx in [0, 4]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(val)
        if c_idx == 4:
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x00, 0x64, 0x00)
        r.font.size = Pt(9.5)

add_sec("5.2. Tỷ lệ hoàn thành công việc")
add_para("Bảng đánh giá tỷ lệ đóng góp của từng thành viên trong nhóm được thể hiện tại Bảng 5.2 dưới đây.")
add_tbl_caption("Bảng 5.2: Bảng đánh giá tỷ lệ đóng góp của từng thành viên trong nhóm")

contrib_tbl = doc.add_table(rows=5, cols=5)
contrib_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(contrib_tbl)

ct_hdrs = ["STT", "Họ và tên thành viên", "Nhiệm vụ đảm nhiệm chính", "Mức độ hoàn thành", "Chữ ký xác nhận"]
for i, h in enumerate(ct_hdrs):
    c = contrib_tbl.rows[0].cells[i]
    set_cell_background(c, "14221D")
    set_cell_margins(c, 50, 50, 60, 60)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

contrib_data = [
    ("1", "Nguyễn Huy Hoàng", "Kiến trúc hệ thống, Backend, CSDL, Triggers và Kiểm thử tự động", "100%", ""),
    ("2", "Nguyễn Việt Vương", "Giao diện Blade, Tương thích di động và Tương tác JavaScript", "100%", ""),
    ("3", "Nguyễn Ngọc Tường Vi", "Chuẩn hóa dữ liệu Seeder, Chức năng Khách hàng và Báo cáo kỹ thuật", "100%", ""),
    ("TỔNG", "Cả nhóm", "Hoàn thành toàn bộ các hạng mục công việc được phân công", "100%", "")
]

for r_idx, (c0, c1, c2, c3, c4) in enumerate(contrib_data, start=1):
    row = contrib_tbl.rows[r_idx]
    bg = "F4F6F8" if r_idx == 4 else ("F9FBFA" if r_idx % 2 == 1 else "FFFFFF")
    for c_idx, val in enumerate([c0, c1, c2, c3, c4]):
        c = row.cells[c_idx]
        set_cell_background(c, bg)
        set_cell_margins(c, 40, 40, 50, 50)
        p = c.paragraphs[0]
        if c_idx in [0, 3]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(val)
        if r_idx == 4 or c_idx == 3:
            r.font.bold = True
        r.font.size = Pt(9.5)

add_para("Về phương thức phối hợp làm việc nhóm, nhóm sử dụng GitHub làm nền tảng quản lý mã nguồn tập trung. Nhánh main được quy định là nhánh mã nguồn ổn định, chỉ tiến hành hợp nhất sau khi các tính năng mới đã được thử nghiệm cục bộ và vượt qua toàn bộ 23 kịch bản kiểm thử tự động của PHPUnit. Mỗi thành viên khi phát triển một module mới đều tạo một nhánh tính năng riêng biệt từ main. Trước khi tạo Pull Request để gộp mã nguồn, thành viên chủ động rebase với nhánh main mới nhất nhằm phát hiện sớm và xử lý các điểm xung đột mã nguồn. Trước khi bắt tay vào lập trình cơ sở dữ liệu, nhóm đã họp trực tiếp để thống nhất bảng quy ước tên bảng, tên cột và kiểu dữ liệu chuẩn của 13 bảng nghiệp vụ, nhờ đó các biến truyền từ Controller sang view Blade luôn ăn khớp hoàn toàn mà không phát sinh lỗi lệch tên trường.")

add_sec("5.3. Các hạn chế hiện tại")
add_para("Trong phạm vi thời gian thực hiện học phần Lập trình mã nguồn mở, nhóm tập trung nguồn lực để hoàn thiện các bài toán nghiệp vụ cốt lõi về biến thể kích cỡ, luồng đặt hàng và kiểm soát tranh chấp tồn kho. Do đó, hệ thống hiện tại vẫn còn một số điểm hạn chế kỹ thuật:")
add_para("Thứ nhất, hệ thống chưa kết nối tự động với các cổng thanh toán điện tử trực tuyến. Các tùy chọn thanh toán hiện mới dừng lại ở mức thanh toán khi nhận hàng (COD) và chuyển khoản ngân hàng thủ công. Với hình thức chuyển khoản, khách hàng phải tự quét mã hoặc chuyển tiền theo số tài khoản hiển thị trên màn hình, sau đó nhân viên bán hàng phải đối soát tài khoản ngân hàng thủ công rồi mới bấm nút 'Xác nhận đã thanh toán' trên trang quản trị. Dự án chưa tích hợp API của các cổng trung gian như VNPay hay MoMo kèm cơ chế phản hồi Webhook IPN để hệ thống tự động cập nhật trạng thái đã thu tiền ngay khi giao dịch hoàn tất.")
add_para("Thứ hai, hệ thống chưa tích hợp dịch vụ gửi email thông báo tự động. Các thông báo đơn hàng hiện mới chỉ hiển thị trực tiếp trên giao diện web và lưu vào tài khoản cá nhân. Ứng dụng chưa cấu hình Laravel Mail qua giao thức SMTP để tự động gửi thư xác nhận hóa đơn kèm tệp PDF về hộp thư của khách hàng sau khi thanh toán, cũng như chưa xây dựng luồng đặt lại mật khẩu qua email xác thực khi người dùng quên mật khẩu.")
add_para("Thứ ba, hiệu năng xử lý của trang catalog sản phẩm chưa được hỗ trợ bởi tầng bộ nhớ đệm (Caching). Toàn bộ các thao tác tải trang chủ, lọc danh mục, tìm kiếm theo tên và tính điểm sao trung bình qua Accessor hiện đều gửi truy vấn trực tiếp xuống hệ quản trị MySQL thông qua Eloquent ORM. Khi lượng người dùng truy cập đồng thời tăng cao, việc liên tục thực thi các câu lệnh SELECT phức tạp kèm nhiều mệnh đề WHERE và JOIN sẽ làm tăng tải cho máy chủ cơ sở dữ liệu và kéo dài thời gian phản hồi trang web.")
add_para("Thứ tư, cấu trúc bảng mã giảm giá chưa lưu trữ thời hạn hiệu lực cụ thể. Bảng MA_GIAM_GIA hiện chỉ có cột số lượt dùng SoLuong và cờ trạng thái TrangThai để kiểm tra tính khả dụng, chưa có cột lưu trữ ngày bắt đầu và ngày kết thúc của voucher. Việc dừng áp dụng mã giảm giá hiện hoàn toàn phụ thuộc vào việc quản trị viên chuyển cờ trạng thái thủ công hoặc khi số lượt dùng giảm về 0.")

add_sec("5.4. Kế hoạch nâng cấp và phát triển mở rộng")
add_para("Dựa trên các hạn chế kỹ thuật nêu trên, nhóm xây dựng kế hoạch nâng cấp hệ thống theo 4 giai đoạn cụ thể. Lộ trình 4 giai đoạn nâng cấp và mở rộng hệ thống được mô tả tại Hình 5.1 dưới đây.")
add_box("SƠ ĐỒ LỘ TRÌNH 4 GIAI ĐOẠN NÂNG CẤP VÀ PHÁT TRIỂN HỆ THỐNG")
add_caption("Hình 5.1: Lộ trình 4 giai đoạn nâng cấp và mở rộng hệ thống Smart Casual")

add_para("Giai đoạn 1 tập trung vào việc tích hợp cổng thanh toán trực tuyến qua VNPay Sandbox. Đây là nhiệm vụ được ưu tiên hàng đầu nhằm hoàn thiện chu trình thanh toán không tiền mặt, giúp loại bỏ khâu kiểm tra sao kê thủ công của nhân viên bán hàng và hạn chế rủi ro đơn hàng ảo. Nhóm sẽ tích hợp bộ thư viện VNPay SDK, sử dụng thuật toán HMAC-SHA512 để kiểm tra chữ ký dữ liệu và xây dựng route tiếp nhận phản hồi IPN (/vnpay_return) để tự động cập nhật trường DaThanhToan = 1 ngay khi khách hàng quét mã QR thành công.")
add_para("Giai đoạn 2 hướng đến việc cải thiện trải nghiệm khách hàng thông qua hệ thống gửi email tự động bằng Laravel Mail kết hợp hàng đợi cơ sở dữ liệu (Laravel Queue). Khi khách hàng đặt đơn thành công, hệ thống đẩy một tác vụ SendOrderConfirmationEmailJob vào hàng đợi nền để xử lý việc kết nối máy chủ SMTP và gửi email xác nhận kèm hóa đơn PDF. Việc đưa tác vụ gửi thư vào hàng đợi nền giúp trang thanh toán của khách phản hồi ngay lập tức mà không bị treo chờ kết nối mạng đến máy chủ thư điện tử.")
add_para("Giai đoạn 3 tập trung tối ưu hóa hiệu năng truy vấn bằng bộ nhớ đệm Redis Cache. Nhóm sẽ cấu hình Redis để lưu trữ tạm thời kết quả truy vấn danh mục sản phẩm và danh sách 4 sản phẩm thịnh hành ngoài trang chủ trong khoảng thời gian 60 phút. Khi quản trị viên thực hiện các thao tác thêm, sửa hoặc ẩn sản phẩm, hệ thống tự động xóa bộ nhớ đệm (cache invalidation) để dữ liệu hiển thị ngoài trang bán hàng luôn đồng nhất với cơ sở dữ liệu.")
add_para("Giai đoạn 4 mở rộng phạm vi nền tảng bằng việc bổ sung trường ngày hiệu lực cho bảng MA_GIAM_GIA và xây dựng hệ thống RESTful API chuẩn mực. Các API tại tiền tố /api/v1/* sẽ trả về dữ liệu định dạng JSON và sử dụng thư viện Laravel Sanctum để cấp mã token xác thực, tạo nền tảng cho việc phát triển ứng dụng di động trên nền Flutter hoặc React Native trong tương lai.")

# ----------------- KẾT LUẬN -----------------
add_chapter("KẾT LUẬN")
add_para("Dự án \"Xây dựng Website Bán hàng Thời trang Smart Casual bằng Laravel Framework và MySQL\" đã hoàn thành các mục tiêu đề ra trong học phần Lập trình mã nguồn mở tại Trường Đại học Công Thương TP. Hồ Chí Minh. Hệ thống được xây dựng trên nền tảng Laravel 11.x kết hợp cơ sở dữ liệu quan hệ MySQL 8.4, đáp ứng trọn vẹn chu trình bán lẻ thời trang với 13 bảng nghiệp vụ chuẩn hóa, từ quản lý danh mục, biến thể kích cỡ sản phẩm, giỏ hàng, quy trình đặt hàng đến bảng điều khiển quản trị phân quyền cho nhân viên và người quản lý.")
add_para("Trong quá trình thực hiện đồ án, nhóm đã đối mặt và giải quyết được hai khó khăn kỹ thuật then chốt: thứ nhất là bài toán tranh chấp dữ liệu khi nhiều khách hàng cùng đặt mua một biến thể sản phẩm chỉ còn tồn kho bằng 1, nhóm đã xử lý bằng cơ chế khóa dòng bi quan lockForUpdate() kết hợp Database Transaction tại lớp dịch vụ CheckoutService; thứ hai là việc Eloquent ORM mặc định không hỗ trợ khóa chính phức hợp hai cột trên bảng GIO_HANG, nhóm đã giải quyết bằng cách ghi đè phương thức setKeysForSaveQuery để can thiệp trực tiếp vào câu lệnh SQL cập nhật.")
add_para("Thông qua đồ án này, các thành viên trong nhóm đã nắm vững phương pháp tổ chức mã nguồn theo mô hình MVC kết hợp Service Layer, hiểu rõ cách thức phân quyền người dùng qua Multi-Auth Guards và Middleware tùy chỉnh, đồng thời làm quen với việc phối hợp kỹ thuật giữa tầng ứng dụng và Database Triggers. Bộ 23 kịch bản kiểm thử tự động được xây dựng cũng giúp nhóm rèn luyện kỹ năng kiểm thử phần mềm có phương pháp, đảm bảo các tính năng vận hành ổn định trước khi đưa vào sử dụng thực tế.")

# ----------------- TÀI LIỆU THAM KHẢO -----------------
add_chapter("TÀI LIỆU THAM KHẢO")
refs = [
    "[1] Khoa Công nghệ Thông tin — Trường Đại học Công Thương TP. HCM, Đề cương chi tiết học phần Lập trình mã nguồn mở, Bộ môn Kỹ thuật Phần mềm, 2026.",
    "[2] T. Otwell, Laravel Documentation: The PHP Framework for Web Artisans (Version 11.x), Laravel LLC, 2026. [Trực tuyến]. Khả dụng: https://laravel.com/docs/11.x",
    "[3] The PHP Group, PHP 8.3 Manual and Language Reference, 2026. [Trực tuyến]. Khả dụng: https://www.php.net/manual/en/",
    "[4] Oracle Corporation, MySQL 8.4 Reference Manual, Oracle Corporation, 2026. [Trực tuyến]. Khả dụng: https://dev.mysql.com/doc/refman/8.4/en/",
    "[5] M. Otto và J. Thornton, Bootstrap 5.2 Documentation: The most popular HTML, CSS, and JS library in the world, Bootstrap Core Team, 2026. [Trực tuyến]. Khả dụng: https://getbootstrap.com/docs/5.2/",
    "[6] S. Bergmann, PHPUnit 10 Manual: The PHP Testing Framework, 2026. [Trực tuyến]. Khả dụng: https://phpunit.de/documentation.html"
]

for r_text in refs:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(r_text)
    r.font.size = Pt(11)

# Save document
output_path1 = Path(r"D:\PHP_NOSQL\BAO_CAO_DO_AN_MON_HOC_LARAVEL_5_CHUONG.docx")
doc.save(str(output_path1))
print(f"Successfully compiled: {output_path1}")

output_path2 = Path(r"D:\PHP_NOSQL\SMART_CASUAL_LARAVEL\docs\BAO_CAO_DO_AN_MON_HOC_LARAVEL_5_CHUONG.docx")
doc.save(str(output_path2))
print(f"Successfully compiled: {output_path2}")
